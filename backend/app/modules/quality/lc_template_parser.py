"""液相计算表模板配置驱动的通用解析器。

模板识别键 = 表号（EX-xx-xxxx-vvv，标题单元格内）。取值位置由 quality_lc_template_configs
的 config JSON 描述，支持两种区块：
- paired_rows：组分名标签列（如 A 列「Vancomycin B％」）+ 每组分固定行距（第一份/第二份），
  结果列（报告值，TEXT 百分比文本或数值），可按配置取两份平均或回落列；
- fixed_cell：固定单元格（如妥布霉素含量 L12）。
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from io import BytesIO

EX_RE = re.compile(r"EX-[A-Z]+-\d+-\d+")
PCT_TEXT_RE = re.compile(r"^\(?([\d.]+)\)?\s*%$")


@dataclass
class ParsedComponent:
    name: str
    first: float | None = None
    second: float | None = None
    report_value: float | None = None  # 最终报告值（百分数数值，如 96.3）
    limit: float | None = None  # 计算表自带的合格限度（百分数数值，如 93.0）


@dataclass
class GenericLcParse:
    product_name: str = ""
    batch_number: str = ""
    form_id: str = ""
    components: list[ParsedComponent] = field(default_factory=list)


def _col_letter(c: int) -> str:
    letters = ""
    while c > 0:
        c, rem = divmod(c - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def _cell_coord(cell: str) -> tuple[int, int]:
    """'L12' → (12, 12)。"""
    m = re.match(r"([A-Z]+)(\d+)", cell.upper())
    if not m:
        raise ValueError(f"非法单元格坐标: {cell}")
    col = 0
    for ch in m.group(1):
        col = col * 26 + ord(ch) - 64
    return int(m.group(2)), col


def _parse_value(v) -> float | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s in ("", "-", "未检出"):
        return None
    m = PCT_TEXT_RE.match(s)
    if m:
        return float(m.group(1))
    try:
        return float(s.replace(",", ""))
    except ValueError:
        return None


def open_sheet(file_bytes: bytes, filename: str):
    """返回 (getter(r, c), max_row, max_col)。xlsx 用 openpyxl（缓存值），xls 用 xlrd。"""
    if filename.lower().endswith(".xls"):
        import xlrd

        book = xlrd.open_workbook(file_contents=file_bytes)
        sh = book.sheet_by_index(0)

        def get(r: int, c: int):
            if r < 1 or c < 1 or r > sh.nrows or c > sh.ncols:
                return None
            return sh.cell_value(r - 1, c - 1)

        return get, sh.nrows, sh.ncols

    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(file_bytes), data_only=True)
    ws = wb.active
    # 合并单元格映射：非锚点格读锚点值（报告值常一格跨多行，如 M22:M25）
    merge_anchor: dict[tuple[int, int], tuple[int, int]] = {}
    for mr in ws.merged_cells.ranges:
        anchor = (mr.min_row, mr.min_col)
        for rr in range(mr.min_row, mr.max_row + 1):
            for cc in range(mr.min_col, mr.max_col + 1):
                merge_anchor[(rr, cc)] = anchor

    def get(r: int, c: int):
        if r < 1 or c < 1 or r > ws.max_row or c > ws.max_column:
            return None
        ar, ac = merge_anchor.get((r, c), (r, c))
        return ws.cell(ar, ac).value

    return get, ws.max_row, ws.max_column


def detect_table_no(get, max_row: int, max_col: int) -> str | None:
    """标题单元格（前 5 行）中提取表号 EX-xx-xxxx-vvv。"""
    for r in range(1, min(6, max_row + 1)):
        for c in range(1, min(12, max_col + 1)):
            v = get(r, c)
            if v and isinstance(v, str):
                m = EX_RE.search(v)
                if m:
                    return m.group(0)
    return None


def find_batch(get, max_row: int, max_col: int, label: str = "批号") -> str:
    """定位「批号」标签：值在同格冒号后，或标签右侧一格。"""
    for r in range(1, min(15, max_row + 1)):
        for c in range(1, min(8, max_col + 1)):
            v = get(r, c)
            if v and isinstance(v, str) and label in v:
                rest = v.replace(label, "").strip().lstrip("：:").strip()
                if rest:
                    return rest
                nxt = get(r, c + 1)
                if nxt is not None and str(nxt).strip():
                    return str(nxt).strip()
    return ""


def _apply_merges(
    components: list[ParsedComponent], merge_rules: list[dict]
) -> list[ParsedComponent]:
    """组分合并规则：如 杂质B = 杂质B1 + 杂质B2（取二者之和）。

    规则形如 {"into": "杂质B(B1+B2)", "from": ["杂质B1", "杂质B2"], "op": "sum"}。
    按名称包含关系收集 from 组分，op=sum 时报告值/份值分别求和，合并为一个组分。
    """
    for rule in merge_rules:
        into = rule.get("into", "")
        parts = rule.get("from", [])
        op = rule.get("op", "sum")
        if not into or not parts:
            continue
        matched = [c for c in components if any(p in c.name for p in parts)]
        if not matched:
            continue
        if op == "sum":
            first = sum(c.first for c in matched if c.first is not None) or None
            second = sum(c.second for c in matched if c.second is not None) or None
            vals = [c.report_value for c in matched if c.report_value is not None]
            report = round(sum(vals), 4) if vals else None
        else:
            continue
        merged = ParsedComponent(name=into, first=first, second=second, report_value=report)
        components = [c for c in components if c not in matched] + [merged]
    return components


def parse_with_config(
    file_bytes: bytes, filename: str, config: dict, form_id: str
) -> GenericLcParse:
    get, max_row, max_col = open_sheet(file_bytes, filename)
    result = GenericLcParse(
        product_name=config.get("product_name", ""),
        batch_number=find_batch(get, max_row, max_col, config.get("batch_label", "批号")),
        form_id=form_id,
    )
    for block in config.get("blocks", []):
        kind = block.get("kind")
        if kind == "fixed_cell":
            r, c = _cell_coord(block["cell"])
            result.components.append(ParsedComponent(
                name=block.get("name", ""), report_value=_parse_value(get(r, c)),
            ))
        elif kind == "paired_rows":
            _, name_col = _cell_coord(block["name_col"] + "1")
            _, val_col = _cell_coord((block.get("value_col") or "A") + "1")
            fallback_col = block.get("fallback_col")
            _, fb_col = _cell_coord((fallback_col or block.get("value_col") or "A") + "1")
            step = int(block.get("step", 4))
            avg = bool(block.get("avg", False))
            r = int(block.get("start_row", 1))
            while r <= max_row:
                label = get(r, name_col)
                # 组分标签带百分比号（模板里全角 ％ 与半角 % 均有）
                matched = label and isinstance(label, str) and (
                    "%" in label or "％" in label
                )
                if not matched:
                    r += 1
                    continue
                name = label.strip()
                half = step // 2

                def pick(row: int) -> float | None:
                    v = _parse_value(get(row, val_col))
                    if v is None and fb_col != val_col:
                        v = _parse_value(get(row, fb_col))
                    return v

                first = pick(r)
                second = pick(r + half)
                if avg:
                    vals = [x for x in (first, second) if x is not None]
                    report = round(sum(vals) / len(vals), 4) if vals else None
                else:
                    report = first if first is not None else second
                result.components.append(ParsedComponent(
                    name=name, first=first, second=second, report_value=report,
                ))
                r += step
    result.components = _apply_merges(result.components, config.get("merge", []))
    _apply_limits(get, max_row, max_col, result.components, config.get("limits"))
    return result


def _apply_limits(
    get, max_row: int, max_col: int,
    components: list[ParsedComponent], limits_cfg: dict | None,
) -> None:
    """计算表自带限度列（如 O 列组分名 + P 列限度，小数形式 ×100 → 百分数）。

    limits_cfg = {"name_col": "O", "value_col": "P", "scale": 100}。
    按名称包含关系把限度挂到对应组分（未匹配到标准库的组分靠它自动追加任务行）。
    """
    if not limits_cfg:
        return
    _, name_col = _cell_coord(limits_cfg.get("name_col", "O") + "1")
    _, value_col = _cell_coord(limits_cfg.get("value_col", "P") + "1")
    scale = float(limits_cfg.get("scale", 100))
    limits: dict[str, float] = {}
    for r in range(1, max_row + 1):
        name = get(r, name_col)
        if not name or not isinstance(name, str):
            continue
        if "注意" in name or "组分" in name or "超趋势" in name:
            continue
        v = _parse_value(get(r, value_col))
        if v is not None:
            limits[name.strip()] = round(v * scale, 6)
    for c in components:
        for name, val in limits.items():
            if name in c.name or c.name.replace("％", "").replace("%", "") in name:
                c.limit = val
                break
