"""危化品 Excel 报表解析器（12 张异构部门 sheet → 归一 ParsedRecord）。

数据源：2026.08.23危险品报表.xls（真实异构数据，字段错位/单位混杂较多）。
策略：每 sheet 一个解析函数，尽量抽取 部门/存放部位/物料/包装/数量/单位/总量T/上限/危险性；
明显脏数据跳过或尽力解析，后续人工在表内修正。
"""
from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Any

import xlrd  # type: ignore[import-untyped]

from app.modules.safety.chemical_inventory.rules import _normalize_to_tonnes

# 每日部门报表 sheet 名 → 解析 key（08.25 起各部门单独发 .xlsx，sheet 名与 08.23 汇总不一致）
_DAILY_SHEET_ALIAS: dict[str, str] = {
    "危化品库存汇总": "仓库",
    "总表": "提炼二期",
    "精制": "精制 ",
    "QC危化品库存": "QC",
    "万古": "提一",
    "发一危化品库存 2026": "发一",
    # 09-02 起菌种改版文件 sheet 名无尾随空格，旧表为「菌种 」
    "菌种": "菌种 ",
}

# sheet 名 → 部门枚举
SHEET_DEPARTMENT: dict[str, str] = {
    "仓库": "warehouse",
    "提一": "extraction_1",
    "提二 ": "extraction_2b",
    "提炼二期": "extraction_2",
    "发一": "fermentation_1",
    "发二": "fermentation_2",
    "精制 ": "purification",
    "菌种 ": "strain",
    "QC": "qc",
    "环保": "env",
    "提炼半合成工程中心": "semi_synth",
    "提炼精进中心": "tech_refine",
}

# 危险性文本 → 枚举
_HAZARD_TEXT_MAP: list[tuple[str, str]] = [
    ("易制毒", "precursor_drug"),
    ("易制爆", "precursor_explosive"),
    ("易燃", "flammable"),
    ("易爆", "explosive"),
    ("氧化", "oxidizer"),
    ("腐蚀", "corrosive"),
    ("毒性", "toxic"),
    ("刺激", "irritant"),
]

_UNIT_MAP: dict[str, str] = {
    "kg": "kg", "Kg": "kg", "KG": "kg", "㎏": "kg", "千克": "kg",
    "g": "g", "G": "g", "克": "g",
    "t": "T", "T": "T", "吨": "T", "ton": "T",
    "l": "L", "L": "L", "升": "L",
    "ml": "ml", "ML": "ml", "毫升": "ml",
    "瓶": "bottle", "桶": "bottle", "包": "bottle", "袋": "bottle",
}


@dataclass
class ParsedRecord:
    department: str
    storage_location: str | None
    material_name: str
    package_spec: str | None = None
    quantity: Decimal | None = None
    unit: str | None = None
    total_quantity_t: Decimal | None = None
    max_limit: Decimal | None = None
    max_limit_unit: str | None = None
    hazard_classes: list[str] = field(default_factory=list)
    category: str | None = None
    remark: str | None = None


def _cell(sh: Any, r: int, c: int) -> Any:
    """安全取单元格（列越界返回空串）。"""
    if c >= sh.ncols:
        return ""
    return sh.cell_value(r, c)


class _XlsxSheet:
    """openpyxl worksheet → 兼容 xlrd 的 .nrows/.ncols/.cell_value。"""

    def __init__(self, ws: Any) -> None:
        self.ws = ws
        self.nrows = ws.max_row or 0
        self.ncols = ws.max_column or 0

    def cell_value(self, r: int, c: int) -> Any:
        if r < 0 or r >= self.nrows or c < 0 or c >= self.ncols:
            return ""
        v = self.ws.cell(row=r + 1, column=c + 1).value
        return "" if v is None else v


def _read_sheets(path: str) -> list[tuple[str, Any]]:
    """读取 .xls（xlrd）或 .xlsx（openpyxl），返回 [(sheet名, sheet对象)]。"""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xlsx":
        import openpyxl  # type: ignore[import-untyped]

        wb = openpyxl.load_workbook(path, data_only=True)
        return [(n, _XlsxSheet(wb[n])) for n in wb.sheetnames]
    wb = xlrd.open_workbook(path)
    return [(n, wb.sheet_by_name(n)) for n in wb.sheet_names()]


def _num(v: Any) -> Decimal | None:
    if v is None:
        return None
    if isinstance(v, Decimal):
        return v
    s = str(v).strip().replace(",", "").replace("，", "")
    if not s or s in ("-", "/", "—", "暂未使用", "无"):
        return None
    try:
        return Decimal(s)
    except (InvalidOperation, ValueError):
        return None


def _norm_unit(u: Any) -> str | None:
    if u is None:
        return None
    s = str(u).strip()
    if not s:
        return None
    return _UNIT_MAP.get(s, s.lower() if len(s) <= 3 else s)


def _hazard_classes(text: Any) -> list[str]:
    if not text:
        return []
    s = str(text)
    result: list[str] = []
    # 「易燃易爆」为口语化标签，指易燃液体（如乙醇/丙酮），并非易爆品；
    # 此处只映射为 flammable，避免把常见溶剂误标为 explosive。
    composite_explosive = "易燃易爆" in s
    for kw, enum in _HAZARD_TEXT_MAP:
        if kw in s and enum not in result:
            if enum == "explosive" and composite_explosive:
                continue
            result.append(enum)
    return result


def _parse_limit(v: Any) -> tuple[Decimal | None, str | None]:
    """从 库存上限/最大存放量 单元格解析 (数值, 单位)。如 '3T' / '4000kg' / '30000.0' / '200㎏'。"""
    if v is None:
        return None, None
    s = str(v).strip().replace("，", "").replace(" ", "")
    if not s or s in ("-", "/", "—", "暂未使用"):
        return None, None
    m = re.match(r"^([0-9.]+)([A-Za-z㎀-㏿一-鿿]*)", s)
    if not m:
        return _num(s), None
    num = _num(m.group(1))
    unit = _norm_unit(m.group(2)) if m.group(2) else None
    return num, unit


# 库存单元格内联单位：'0ml' / '490g' / '0.5L' / '2吨' / 纯数字
_QTY_RE = re.compile(r"^([0-9][0-9,，.]*)\s*([A-Za-z㎀-㏿一-鿿]{0,3})?$")


def _parse_qty(v: Any) -> tuple[Decimal | None, str | None]:
    """解析库存单元格，返回 (数值, 内联单位)。

    菌种等部门的日报把单位直接写在库存里（如 '0ml'、'490g'），纯 _num 解析会失败；
    内联单位仅作用于库存本身，上限的默认单位以「单位」列为准。
    """
    if v is None:
        return None, None
    s = str(v).strip()
    if not s:
        return None, None
    m = _QTY_RE.match(s)
    if not m:
        return _num(s), None
    num = _num(m.group(1))
    unit = _norm_unit(m.group(2)) if m.group(2) else None
    return num, unit


def _limit_to_tonnes(raw: Any, default_unit: str | None) -> tuple[Decimal | None, str | None]:
    """把库存上限统一换算成 T（规则 R01-R03 需与总量同单位）。"""
    val, u = _parse_limit(raw)
    if val is None:
        return None, None
    uu = u or default_unit
    if not uu:
        return val, None
    t = _normalize_to_tonnes(val, uu, None)
    if t is None:
        return val, uu
    return t, "T"


def _parse_package_spec(v: Any) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    if not s or s in ("-", "/", "—"):
        return None
    return s


def _clean_name(v: Any) -> str:
    s = str(v).strip().replace("\n", "").replace("\n", "")
    return s


# ── 各 sheet 解析 ──


def _parse_warehouse(sh: xlrd.sheet.Sheet, dept: str) -> list[ParsedRecord]:
    """仓库：库位|物料名称|包装规格|库存数量|单位|Kg|储存位置|最大存放量|现场物料总量。"""
    out: list[ParsedRecord] = []
    for r in range(2, sh.nrows):
        loc = _clean_name(sh.cell_value(r, 0))
        name = _clean_name(sh.cell_value(r, 1))
        spec = _parse_package_spec(sh.cell_value(r, 2))
        qty = _num(sh.cell_value(r, 3))
        unit = _norm_unit(sh.cell_value(r, 4))
        limit, limit_unit = _limit_to_tonnes(sh.cell_value(r, 7), "T")
        if not name or qty is None or qty == 0 or "合计" in name:
            continue
        total = _normalize_to_tonnes(qty, unit, spec)
        out.append(ParsedRecord(
            department=dept, storage_location=loc or None, material_name=name,
            package_spec=spec, quantity=qty, unit=unit, total_quantity_t=total,
            max_limit=limit, max_limit_unit=limit_unit,
        ))
    return out


def _parse_tiyi(sh: xlrd.sheet.Sheet, dept: str) -> list[ParsedRecord]:
    """提一：序号|存放部位|危险品名称|规格|车间地面/kg|车间罐内/kg|未折纯量/kg|折纯量/kg|库存上限/kg|危险性。"""
    out: list[ParsedRecord] = []
    for r in range(2, sh.nrows):
        name = _clean_name(sh.cell_value(r, 2))
        if not name or "汇总" in name or "品类" in name or "求和项" in name:
            continue
        if _num(name) is not None:
            continue  # 汇总段的数值行
        # 汇总段（序号列为空）跳过
        if not str(sh.cell_value(r, 0)).strip():
            continue
        loc = _clean_name(sh.cell_value(r, 1)) or None
        spec = _parse_package_spec(sh.cell_value(r, 3))
        qty = _num(sh.cell_value(r, 6))  # 未折纯量/kg
        if qty is None:
            ground = _num(sh.cell_value(r, 4)) or Decimal("0")
            tank = _num(sh.cell_value(r, 5)) or Decimal("0")
            qty = ground + tank
        limit_kg = _num(sh.cell_value(r, 8))
        limit = limit_kg / 1000 if limit_kg is not None else None
        hazard = _hazard_classes(sh.cell_value(r, 9))
        total = qty / 1000 if qty is not None else None  # kg → T
        out.append(ParsedRecord(
            department=dept, storage_location=loc, material_name=name,
            package_spec=spec, quantity=qty, unit="kg", total_quantity_t=total,
            max_limit=limit, max_limit_unit="T", hazard_classes=hazard,
        ))
    return out


def _parse_tier(sh: xlrd.sheet.Sheet, dept: str) -> list[ParsedRecord]:
    """提炼二期：序号|车间|存放部位|危险品名称|规格|未折纯量/kg|折纯量/kg|库存上限|品类|品种。"""
    out: list[ParsedRecord] = []
    for r in range(2, sh.nrows):
        name = _clean_name(sh.cell_value(r, 3))
        if not name:
            continue
        loc = _clean_name(sh.cell_value(r, 2)) or None
        spec = _parse_package_spec(sh.cell_value(r, 4))
        qty = _num(sh.cell_value(r, 5))  # 未折纯量/kg
        limit, limit_unit = _limit_to_tonnes(sh.cell_value(r, 7), "kg")
        category = _clean_name(sh.cell_value(r, 8)) or None
        total = qty / 1000 if qty is not None else None
        out.append(ParsedRecord(
            department=dept, storage_location=loc, material_name=name,
            package_spec=spec, quantity=qty, unit="kg", total_quantity_t=total,
            max_limit=limit, max_limit_unit=limit_unit or "T", category=category,
        ))
    return out


def _parse_tierbu(sh: xlrd.sheet.Sheet, dept: str) -> list[ParsedRecord]:
    """提二（提炼二部）：列错位严重，尽力抽 存放部位|物料名称|包装规格|库存|上限|单位。"""
    out: list[ParsedRecord] = []
    cur_loc = None
    for r in range(2, sh.nrows):
        loc = _clean_name(sh.cell_value(r, 1))
        if loc:
            cur_loc = loc
        name = _clean_name(sh.cell_value(r, 2))
        if not name:
            continue
        spec = _parse_package_spec(sh.cell_value(r, 3))
        qty = _num(sh.cell_value(r, 5))  # 库存量（kg/L/t）
        if qty is None:
            qty = _num(sh.cell_value(r, 4))  # 库存（包/瓶数）
        unit = _norm_unit(sh.cell_value(r, 7))
        limit, limit_unit = _limit_to_tonnes(sh.cell_value(r, 6), unit)
        hazard = _hazard_classes(sh.cell_value(r, 9))
        total = _normalize_to_tonnes(qty, unit, spec) if qty is not None else None
        if unit == "t" or unit == "T":
            total = qty  # 已是吨
        out.append(ParsedRecord(
            department=dept, storage_location=cur_loc, material_name=name,
            package_spec=spec, quantity=qty, unit=unit, total_quantity_t=total,
            max_limit=limit, max_limit_unit=limit_unit, hazard_classes=hazard,
        ))
    return out


def _detect_fayifaer_columns(sh: Any) -> tuple[int, dict[str, int], list[int]] | None:
    """前 4 行内找含「物料名称」的表头行，按关键词动态定位列。

    各部门表列序常有调整（菌种 2026-09 起改版为 部门|存放部位|物料名称|包装规格|
    库存|库存上限|单位|备注，库存还带内联单位），表头映射比固定列位稳；
    name/qty/limit 齐全才算命中。返回 (表头行号, 列映射, 库存列候选)。
    """
    for r in range(0, min(4, sh.nrows)):
        cols: dict[str, int] = {}
        qty_cols: list[int] = []
        for c in range(sh.ncols):
            h = str(sh.cell_value(r, c)).strip()
            if not h:
                continue
            # 「库存上限」必须先于「库存」判断，否则会误入库存列
            if "上限" in h or "最大" in h:
                cols.setdefault("limit", c)
            elif "库存" in h:
                qty_cols.append(c)
            elif "物料名称" in h:
                cols.setdefault("name", c)
            elif "存放部位" in h or "存放位置" in h:
                cols.setdefault("loc", c)
            elif "包装规格" in h:
                cols.setdefault("spec", c)
            elif "单位" in h:
                cols.setdefault("unit", c)
            elif "备注" in h:
                cols.setdefault("remark", c)
            elif "危险" in h:
                cols.setdefault("hazard", c)
        if "name" in cols and "limit" in cols and qty_cols:
            cols["qty"] = qty_cols[-1]
            return r, cols, qty_cols
    return None


def _parse_fayifaer_by_header(
    sh: Any, dept: str, detected: tuple[int, dict[str, int], list[int]]
) -> list[ParsedRecord]:
    """按表头映射解析（_parse_fayifaer 的表头分支，供 sheet 名兜底复用）。"""
    header_row, cols, qty_cols = detected
    out: list[ParsedRecord] = []
    cur_loc = None
    for r in range(header_row + 1, sh.nrows):
        loc = _clean_name(_cell(sh, r, cols["loc"])) if "loc" in cols else ""
        if loc:
            cur_loc = loc
        name = _clean_name(_cell(sh, r, cols["name"]))
        if not name:
            continue
        spec = _parse_package_spec(_cell(sh, r, cols["spec"])) if "spec" in cols else None
        # 多个库存列候选时从最右列起取（旧表 库存量 位于 瓶数列 右侧）
        qty = None
        inline_unit = None
        for c in reversed(qty_cols):
            qty, inline_unit = _parse_qty(_cell(sh, r, c))
            if qty is not None:
                break
        # 「单位」列是上限的默认单位；库存内联单位（如 0ml）只作用于库存本身
        col_unit = _norm_unit(_cell(sh, r, cols["unit"])) if "unit" in cols else None
        unit = inline_unit or col_unit
        limit, limit_unit = _limit_to_tonnes(_cell(sh, r, cols["limit"]), col_unit)
        hazard = _hazard_classes(_cell(sh, r, cols["remark"])) if "remark" in cols else []
        if not hazard and "hazard" in cols:
            hazard = _hazard_classes(_cell(sh, r, cols["hazard"]))
        total = _normalize_to_tonnes(qty, unit, spec) if qty is not None else None
        out.append(ParsedRecord(
            department=dept, storage_location=cur_loc, material_name=name,
            package_spec=spec, quantity=qty, unit=unit, total_quantity_t=total,
            max_limit=limit, max_limit_unit=limit_unit, hazard_classes=hazard,
        ))
    return out


def _parse_fayifaer(sh: xlrd.sheet.Sheet, dept: str) -> list[ParsedRecord]:
    """发一/发二/菌种/环保/精制 相似：部门|存放部位|物料名称|包装规格|库存|...|库存上限|单位|备注|危险性。

    优先按表头关键词定位列；无表头时回退 08.23 汇总版固定列位（_parse_fayifaer_legacy）。
    """
    detected = _detect_fayifaer_columns(sh)
    if detected is None:
        return _parse_fayifaer_legacy(sh, dept)
    return _parse_fayifaer_by_header(sh, dept, detected)


def _parse_fayifaer_legacy(sh: xlrd.sheet.Sheet, dept: str) -> list[ParsedRecord]:
    """无表头回退：固定列位 库存量 col5（发二/菌种为 col4）、上限 col6、单位 col7。"""
    out: list[ParsedRecord] = []
    cur_loc = None
    for r in range(2, sh.nrows):
        loc = _clean_name(sh.cell_value(r, 1))
        if loc:
            cur_loc = loc
        name = _clean_name(sh.cell_value(r, 2))
        if not name:
            continue
        spec = _parse_package_spec(sh.cell_value(r, 3))
        qty, inline_unit = _parse_qty(sh.cell_value(r, 5))  # 库存量（发一/环保 col5；发二/菌种 col5 为空）
        if qty is None:
            qty, inline_unit = _parse_qty(sh.cell_value(r, 4))  # 发二/菌种：col4 即库存量
        col_unit = _norm_unit(sh.cell_value(r, 7))
        unit = inline_unit or col_unit
        limit, limit_unit = _limit_to_tonnes(sh.cell_value(r, 6), col_unit)
        hazard = _hazard_classes(_cell(sh, r, 8)) or _hazard_classes(_cell(sh, r, 9))
        total = _normalize_to_tonnes(qty, unit, spec) if qty is not None else None
        out.append(ParsedRecord(
            department=dept, storage_location=cur_loc, material_name=name,
            package_spec=spec, quantity=qty, unit=unit, total_quantity_t=total,
            max_limit=limit, max_limit_unit=limit_unit, hazard_classes=hazard,
        ))
    return out


def _parse_jingzhi(sh: xlrd.sheet.Sheet, dept: str) -> list[ParsedRecord]:
    """精制：存放部位|物料名称|包装规格|库存|现场|单位|最大库存|危险性。"""
    out: list[ParsedRecord] = []
    cur_loc = None
    for r in range(2, sh.nrows):
        loc = _clean_name(sh.cell_value(r, 0))
        if loc:
            cur_loc = loc
        name = _clean_name(sh.cell_value(r, 1))
        if not name:
            continue
        spec = _parse_package_spec(sh.cell_value(r, 2))
        qty = _num(sh.cell_value(r, 4))
        unit = _norm_unit(sh.cell_value(r, 5)) or "L"
        if qty is None:
            qty = _num(sh.cell_value(r, 3))
        limit, limit_unit = _limit_to_tonnes(sh.cell_value(r, 6), unit)
        hazard = _hazard_classes(sh.cell_value(r, 7))
        total = _normalize_to_tonnes(qty, unit, spec) if qty is not None else None
        out.append(ParsedRecord(
            department=dept, storage_location=cur_loc, material_name=name,
            package_spec=spec, quantity=qty, unit=unit, total_quantity_t=total,
            max_limit=limit, max_limit_unit=limit_unit, hazard_classes=hazard,
        ))
    return out


def _parse_qc(sh: xlrd.sheet.Sheet, dept: str) -> list[ParsedRecord]:
    """QC：存放位置|试剂名称|包装规格|库存（瓶）|库存（L或KG）|库存上限（L或KG）|单位。"""
    out: list[ParsedRecord] = []
    cur_loc = None
    for r in range(2, sh.nrows):
        loc = _clean_name(sh.cell_value(r, 0))
        if loc:
            cur_loc = loc
        name = _clean_name(sh.cell_value(r, 1))
        if not name:
            continue
        spec = _parse_package_spec(sh.cell_value(r, 2))
        unit = _norm_unit(sh.cell_value(r, 6)) or "L"
        qty = _num(sh.cell_value(r, 4))  # 库存（L或KG）
        if qty is None:
            qty = _num(sh.cell_value(r, 3))  # 库存（瓶）
        limit, limit_unit = _limit_to_tonnes(sh.cell_value(r, 5), unit)
        total = _normalize_to_tonnes(qty, unit, spec) if qty is not None else None
        out.append(ParsedRecord(
            department=dept, storage_location=cur_loc, material_name=name,
            package_spec=spec, quantity=qty, unit=unit, total_quantity_t=total,
            max_limit=limit, max_limit_unit=limit_unit,
        ))
    return out


def _parse_bhsj(sh: xlrd.sheet.Sheet, dept: str) -> list[ParsedRecord]:
    """提炼半合成/提炼精进：存放部位|物料名称|包装规格|库存(瓶)|库存量|单位|最大库存量。

    两 sheet 的「单位」与「上限」列序相反：
    - 提炼半合成工程中心：col5=单位, col6=最大库存量
    - 提炼精进中心：      col5=库存上限, col6=单位
    """
    out: list[ParsedRecord] = []
    cur_loc = None
    for r in range(2, sh.nrows):
        loc = _clean_name(sh.cell_value(r, 0))
        if loc:
            cur_loc = loc.replace("\n", "/")
        name = _clean_name(sh.cell_value(r, 1))
        if not name:
            continue
        spec = _parse_package_spec(sh.cell_value(r, 2))
        qty = _num(sh.cell_value(r, 4))  # 库存量（ml/L/g/kg）
        if qty is None:
            qty = _num(sh.cell_value(r, 3))  # 库存（瓶）
        if dept == "semi_synth":
            unit = _norm_unit(sh.cell_value(r, 5))
            limit_raw = sh.cell_value(r, 6)
        else:  # 提炼精进中心
            unit = _norm_unit(sh.cell_value(r, 6))
            limit_raw = sh.cell_value(r, 5)
        limit, limit_unit = _limit_to_tonnes(limit_raw, unit)
        total = _normalize_to_tonnes(qty, unit, spec) if qty is not None else None
        out.append(ParsedRecord(
            department=dept, storage_location=cur_loc, material_name=name,
            package_spec=spec, quantity=qty, unit=unit, total_quantity_t=total,
            max_limit=limit, max_limit_unit=limit_unit,
        ))
    return out


_PARSERS = {
    "仓库": _parse_warehouse,
    "提一": _parse_tiyi,
    "提二 ": _parse_tierbu,
    "提炼二期": _parse_tier,
    "发一": _parse_fayifaer,
    "发二": _parse_fayifaer,
    "精制 ": _parse_jingzhi,
    "菌种 ": _parse_fayifaer,
    "环保": _parse_fayifaer,
    "QC": _parse_qc,
    "提炼半合成工程中心": _parse_bhsj,
    "提炼精进中心": _parse_bhsj,
}


def parse_workbook(path: str) -> list[ParsedRecord]:
    """解析整本 Excel（跳过无解析器的 sheet），返回归一记录列表。

    支持 .xls（xlrd）与 .xlsx（openpyxl），并通过 _DAILY_SHEET_ALIAS 兼容
    每日部门单表（sheet 名与 08.23 汇总不一致）的场景。
    """
    records: list[ParsedRecord] = []
    for name, sh in _read_sheets(path):
        canonical = _DAILY_SHEET_ALIAS.get(name, name)
        dept = SHEET_DEPARTMENT.get(canonical)
        parser = _PARSERS.get(canonical)
        if not dept or not parser:
            continue
        records.extend(parser(sh, dept))
    return records
