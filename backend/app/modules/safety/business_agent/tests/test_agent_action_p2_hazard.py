"""Agent 执行面二期（agent-action-p2）隐患/报表动作工具测试。

覆盖 spec（.scratch/agent-action-p2/spec.md）T1-T3：
send_hazard_urge 点对点催办 / export_hazard_report 报表导出 /
push_report_to_chat 通用推群。

测试接缝（与一期 test_agent_query_p1_tools 一致）：
- 工具函数边界直接调用（不走 core.pipeline，避免真实审计写库）；
- 依赖层 mock：bitable 链 monkeypatch 源头模块属性、催办核心链 mock
  progress_dunning.send_urge_cards、文件发送 monkeypatch send_file_to_chat；
- 确认流（挂起→确认/取消）复用 tests/modules/safety/test_agent_delegation
  的 executor.resume 打桩模式；
- 路由断言走 tool_selector.select_tool_names 纯函数。
"""

from __future__ import annotations

import base64
import json
from types import SimpleNamespace
from typing import Any

import pytest

from app.modules.safety.business_agent import tool_selector
from app.modules.safety.business_agent.core.permissions import (
    ROLE_DEPT_LEADER,
    ROLE_INSPECTOR,
    ROLE_SAFETY_ADMIN,
    ROLE_VIEWER,
    allowed_tools,
)
from app.modules.safety.business_agent.tools import registry

URGE_ADMIN_LEADER_TOOLS = ["send_hazard_urge"]


# ═══════════════════════════════════════════════════════════════════
# T2 隐患报表导出 — export_hazard_report
# ═══════════════════════════════════════════════════════════════════


class TestExportHazardReportWiring:
    def test_registry_and_groups(self) -> None:
        from app.modules.safety.business_agent.tools.hazard_action_tools import (
            export_hazard_report,
        )

        assert registry.TOOL_KIND["export_hazard_report"] is True
        assert registry._FUNC_BY_NAME["export_hazard_report"] is export_hazard_report
        assert "export_hazard_report" not in registry.RESIDENT_TOOLS
        assert "export_hazard_report" in tool_selector.TOOL_GROUPS["hazard"]

    def test_permissions_admin_and_leader_only(self) -> None:
        assert "export_hazard_report" in allowed_tools(ROLE_SAFETY_ADMIN)
        assert "export_hazard_report" in allowed_tools(ROLE_DEPT_LEADER)
        assert "export_hazard_report" not in allowed_tools(ROLE_INSPECTOR)
        assert "export_hazard_report" not in allowed_tools(ROLE_VIEWER)

    @pytest.mark.parametrize(
        ("message", "expected"),
        [
            ("导一份 7 月的隐患报表", "export_hazard_report"),
            ("导出生产部的隐患清单", "export_hazard_report"),
            ("把隐患台账导出成 Excel", "export_hazard_report"),
        ],
    )
    def test_route_selects_export_tool(self, message: str, expected: str) -> None:
        names = tool_selector.select_tool_names(message)
        assert expected in names, f"{message!r} → {names}"


def _patch_export_chain(
    monkeypatch: pytest.MonkeyPatch,
    *,
    views: list[Any] | None = None,
    total: int | None = None,
    external: list[Any] | None = None,
    send_ok: bool = True,
) -> dict[str, Any]:
    """mock 三源拉取与文件发送；捕获发送的字节供 Excel 解包断言。"""
    captured: dict[str, Any] = {"query": [], "send": [], "_last_bytes": b""}

    async def _fake_query(client: Any, **kwargs: Any):
        captured["query"].append(kwargs)
        return (views or []), total if total is not None else len(views or [])

    async def _fake_external() -> list[Any] | None:
        return external

    async def _fake_send(chat_id: str, file_bytes: bytes, file_name: str, **kw: Any):
        captured["send"].append(
            {"chat_id": chat_id, "size": len(file_bytes), "name": file_name},
        )
        captured["_last_bytes"] = file_bytes
        return send_ok

    monkeypatch.setattr(
        "app.modules.safety.service.hazard_direct.bitable_repo.query_hazards",
        _fake_query,
    )
    monkeypatch.setattr(
        "app.modules.safety.service.hazard_direct.external_audit.fetch_external_audit",
        _fake_external,
    )
    monkeypatch.setattr(
        "app.modules.safety.feishu.chat_sender.send_file_to_chat", _fake_send,
    )
    return captured


def _ext_export_view(
    record_id: str, *, closed: bool = False, desc: str = "罐区围堰破损",
    departments: list[str] | None = None,
    checked_at: Any = None,
) -> Any:
    from datetime import date

    from app.modules.safety.service.hazard_direct.external_audit import (
        ExternalAuditView,
    )

    return ExternalAuditView(
        record_id=record_id, seq=1, description=desc, plan="限期修复",
        departments=departments or ["生产部"], closed=closed,
        in_progress=False, checked_at=checked_at or date(2026, 7, 10),
        target_date=date(2026, 8, 10), responsible=[("张三", "ou_zs")],
    )


def _read_sheet_xlsx(file_bytes: bytes) -> list[list[Any]]:
    import io

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(file_bytes))
    ws = wb.active
    return [[c.value for c in row] for row in ws.iter_rows()]


def _dt(y: int, m: int, d: int) -> Any:
    from datetime import date

    return date(y, m, d)


class TestExportHazardReport:
    def test_internal_export_sends_excel(self, monkeypatch: pytest.MonkeyPatch) -> None:
        from app.modules.safety.business_agent.tools.hazard_action_tools import (
            export_hazard_report,
        )

        captured = _patch_export_chain(
            monkeypatch, views=[_urge_view("recA"), _urge_view("recB")],
        )

        result = _run(export_hazard_report(
            _ctx(None), department="生产部",
            date_from="2026-07-01", date_to="2026-07-31",
        ))

        assert result["success"] is True
        assert result["total"] == 2
        assert result["sent"] is True
        # 文件已发到当前对话，xlsx 后缀
        assert captured["send"], "应调用 send_file_to_chat"
        sent = captured["send"][0]
        assert sent["chat_id"] == "oc_test"
        assert sent["name"].endswith(".xlsx")
        assert sent["size"] > 0
        # 筛选参数进了拉取调用
        q = captured["query"][0]
        assert q.get("department") == "生产部"
        assert q.get("date_from") == "2026-07-01"

    def test_excel_rows_content(self, monkeypatch: pytest.MonkeyPatch) -> None:
        """解包 Excel 验证列头与行内容（外部行为：拿到的文件里有正确数据）。"""
        from app.modules.safety.business_agent.tools.hazard_action_tools import (
            export_hazard_report,
        )

        view = _urge_view("recA", no="HZ-2026-0701")
        view.discovered_at = _dt(2026, 7, 2)
        view.rectification_status = "pending"
        captured = _patch_export_chain(monkeypatch, views=[view])

        _run(export_hazard_report(_ctx(None)))

        rows = _read_sheet_xlsx(captured["_last_bytes"])
        header = rows[0]
        assert header[:3] == ["隐患编号", "来源", "部门"]
        row_text = "".join(str(c) for c in rows[1])
        assert "HZ-2026-0701" in row_text
        assert "内部" in row_text
        assert "未关闭" in row_text

    def test_external_source_export(self, monkeypatch: pytest.MonkeyPatch) -> None:
        from app.modules.safety.business_agent.tools.hazard_action_tools import (
            export_hazard_report,
        )

        captured = _patch_export_chain(
            monkeypatch, external=[_ext_export_view("e1"), _ext_export_view(
                "e2", closed=True, desc="标识缺失",
            )],
        )

        result = _run(export_hazard_report(_ctx(None), source="external"))

        assert result["success"] is True
        assert result["total"] == 2
        rows = _read_sheet_xlsx(captured["_last_bytes"])
        assert len(rows) == 3  # 表头 + 2 行
        assert "外审" in "".join(str(c) for c in rows[1])

    def test_all_source_merges(self, monkeypatch: pytest.MonkeyPatch) -> None:
        from app.modules.safety.business_agent.tools.hazard_action_tools import (
            export_hazard_report,
        )

        _patch_export_chain(
            monkeypatch,
            views=[_urge_view("recA")],
            external=[_ext_export_view("e1")],
        )

        result = _run(export_hazard_report(_ctx(None), source="all"))

        assert result["success"] is True
        assert result["total"] == 2

    def test_over_limit_truncated_hint(
        self, monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        from app.modules.safety.business_agent.tools.hazard_action_tools import (
            export_hazard_report,
        )

        # 真实链路 limit=5000 截断：返回部分行、total 为精确命中数
        captured = _patch_export_chain(
            monkeypatch, views=[_urge_view("recA")], total=5001,
        )

        result = _run(export_hazard_report(_ctx(None)))

        assert result["success"] is True
        assert result["truncated"] is True
        assert result["total"] == 5001
        assert "5000" in result["message"]
        assert captured["send"], "截断时仍应导出已拉到的行"

    def test_zero_hit_returns_error(self, monkeypatch: pytest.MonkeyPatch) -> None:
        from app.modules.safety.business_agent.tools.hazard_action_tools import (
            export_hazard_report,
        )

        captured = _patch_export_chain(monkeypatch, views=[], total=0)

        result = _run(export_hazard_report(_ctx(None), keyword="不存在"))

        assert result["success"] is False
        assert captured["send"] == []

    def test_web_channel_base64_fallback(
        self, monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        from app.modules.safety.business_agent.tools.hazard_action_tools import (
            export_hazard_report,
        )

        _patch_export_chain(monkeypatch, views=[_urge_view("recA")])

        result = _run(export_hazard_report(_ctx(None, chat_id=None)))

        assert result["degraded"] is True
        decoded = base64.b64decode(result["file_content_base64"])
        rows = _read_sheet_xlsx(decoded)
        assert len(rows) == 2

    def test_invalid_source_hint(self) -> None:
        from app.modules.safety.business_agent.tools.hazard_action_tools import (
            export_hazard_report,
        )

        result = _run(export_hazard_report(_ctx(None), source="moon"))

        assert result["success"] is False
        assert "source" in result["error"]

    def test_external_over_limit_truncated(
        self, monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        """外审过滤命中超 5000 → truncated=True 并提示（审查修复回归锚）。"""
        from app.modules.safety.business_agent.tools.hazard_action_tools import (
            export_hazard_report,
        )

        many = [
            _ext_export_view(f"e{i}", desc=f"缺陷{i}") for i in range(5001)
        ]
        captured = _patch_export_chain(monkeypatch, external=many)

        result = _run(export_hazard_report(_ctx(None), source="external"))

        assert result["success"] is True
        assert result["truncated"] is True
        assert result["total"] == 5001
        assert "5000" in result["message"]
        assert captured["send"]

    def test_send_failure_returns_base64(
        self, monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        from app.modules.safety.business_agent.tools.hazard_action_tools import (
            export_hazard_report,
        )

        _patch_export_chain(
            monkeypatch, views=[_urge_view("recA")], send_ok=False,
        )

        result = _run(export_hazard_report(_ctx(None)))

        assert result["sent"] is False
        assert result["file_content_base64"]


# ═══════════════════════════════════════════════════════════════════
# T1 点对点催办 — send_hazard_urge
# ═══════════════════════════════════════════════════════════════════


class TestSendHazardUrgeWiring:
    def test_registry_and_groups(self) -> None:
        from app.modules.safety.business_agent.tools.hazard_action_tools import (
            send_hazard_urge,
        )

        assert registry.TOOL_KIND["send_hazard_urge"] is True
        assert registry._FUNC_BY_NAME["send_hazard_urge"] is send_hazard_urge
        assert "send_hazard_urge" not in registry.RESIDENT_TOOLS
        assert "send_hazard_urge" in tool_selector.TOOL_GROUPS["hazard"]

    def test_permissions_admin_and_leader_only(self) -> None:
        assert "send_hazard_urge" in allowed_tools(ROLE_SAFETY_ADMIN)
        assert "send_hazard_urge" in allowed_tools(ROLE_DEPT_LEADER)
        assert "send_hazard_urge" not in allowed_tools(ROLE_INSPECTOR)
        assert "send_hazard_urge" not in allowed_tools(ROLE_VIEWER)

    @pytest.mark.parametrize(
        ("message", "expected"),
        [
            ("催一下责任人", "send_hazard_urge"),
            ("这条隐患没整改，督办到人", "send_hazard_urge"),
            ("把生产部未关闭的隐患都催一遍", "send_hazard_urge"),
            ("催办张三的隐患整改", "send_hazard_urge"),
        ],
    )
    def test_route_selects_urge_tool(self, message: str, expected: str) -> None:
        names = tool_selector.select_tool_names(message)
        assert expected in names, f"{message!r} → {names}"


def _urge_view(record_id: str, *, no: str = "HZ-2026-001",
               dept: str = "生产部") -> Any:
    from app.modules.safety.service.hazard_direct.bitable_repo import HazardView

    return HazardView(
        record_id=record_id, hazard_no=no,
        description="防爆堵头缺失", department=dept,
        rectification_responsible_person_name="张三",
    )


def _patch_urge_chain(
    monkeypatch: pytest.MonkeyPatch,
    *,
    views: list[Any],
    total: int | None = None,
    urge_result: dict[str, Any] | None = None,
) -> dict[str, list[dict[str, Any]]]:
    """mock 拉取链（bitable_repo.query_hazards）与催办核心链
    （progress_dunning.send_urge_cards）；返回捕获的调用参数。"""

    captured: dict[str, list[dict[str, Any]]] = {"query": [], "urge": []}

    async def _fake_query(client: Any, **kwargs: Any):
        captured["query"].append(kwargs)
        return views, total if total is not None else len(views)

    async def _fake_urge(vws: list[Any]):
        captured["urge"].append({"record_ids": [v.record_id for v in vws]})
        base = (
            urge_result
            if urge_result is not None
            else {
                "total": len(vws), "recipients": 2, "sent": 2,
                "skipped": 0, "errors": 0,
            }
        )
        return base, [], [], []

    monkeypatch.setattr(
        "app.modules.safety.service.hazard_direct.bitable_repo.query_hazards",
        _fake_query,
    )
    monkeypatch.setattr(
        "app.modules.safety.service.hazard_direct.progress_dunning.send_urge_cards",
        _fake_urge,
    )
    return captured


class TestSendHazardUrge:
    def test_single_record_passes_matched_view(self, monkeypatch: pytest.MonkeyPatch) -> None:
        from app.modules.safety.business_agent.tools.hazard_action_tools import (
            send_hazard_urge,
        )

        views = [_urge_view("recA"), _urge_view("recB", no="HZ-2026-002")]
        captured = _patch_urge_chain(monkeypatch, views=views)

        result = _run(send_hazard_urge(_ctx(None), record_id="recA"))

        assert result["success"] is True
        assert result["sent"] == 2
        # 单条入口只把匹配的那条交给催办链
        assert captured["urge"][0]["record_ids"] == ["recA"]
        # 拉取带未关闭语义
        assert captured["query"][0].get("status") == "open"

    def test_single_record_not_found(self, monkeypatch: pytest.MonkeyPatch) -> None:
        from app.modules.safety.business_agent.tools.hazard_action_tools import (
            send_hazard_urge,
        )

        captured = _patch_urge_chain(monkeypatch, views=[_urge_view("recA")])

        result = _run(send_hazard_urge(_ctx(None), record_id="recX"))

        assert result["success"] is False
        assert "recX" in result["error"]
        # 未命中时不应触发催办
        assert captured["urge"] == []

    def test_batch_filtered_urge(self, monkeypatch: pytest.MonkeyPatch) -> None:
        from app.modules.safety.business_agent.tools.hazard_action_tools import (
            send_hazard_urge,
        )

        views = [_urge_view(f"rec{i}", no=f"HZ-2026-0{i}") for i in range(3)]
        captured = _patch_urge_chain(monkeypatch, views=views, total=3)

        result = _run(
            send_hazard_urge(
                _ctx(None), department="生产部",
                date_from="2026-07-01", date_to="2026-07-31",
            )
        )

        assert result["success"] is True
        assert captured["urge"][0]["record_ids"] == ["rec0", "rec1", "rec2"]
        q = captured["query"][0]
        assert q.get("department") == "生产部"
        assert q.get("date_from") == "2026-07-01"

    def test_batch_over_limit_rejected(self, monkeypatch: pytest.MonkeyPatch) -> None:
        from app.modules.safety.business_agent.tools.hazard_action_tools import (
            send_hazard_urge,
        )

        captured = _patch_urge_chain(monkeypatch, views=[], total=51)

        result = _run(send_hazard_urge(_ctx(None), department="生产部"))

        assert result["success"] is False
        assert "50" in result["error"]
        assert captured["urge"] == []

    def test_batch_requires_filter(self, monkeypatch: pytest.MonkeyPatch) -> None:
        from app.modules.safety.business_agent.tools.hazard_action_tools import (
            send_hazard_urge,
        )

        captured = _patch_urge_chain(monkeypatch, views=[])

        result = _run(send_hazard_urge(_ctx(None)))

        assert result["success"] is False
        assert "筛选" in result["error"]
        assert captured["query"] == []

    def test_batch_zero_hit(self, monkeypatch: pytest.MonkeyPatch) -> None:
        from app.modules.safety.business_agent.tools.hazard_action_tools import (
            send_hazard_urge,
        )

        captured = _patch_urge_chain(monkeypatch, views=[], total=0)

        result = _run(send_hazard_urge(_ctx(None), keyword="不存在的关键词"))

        assert result["success"] is True
        assert result["sent"] == 0
        assert captured["urge"] == []

    def test_urge_stats_propagated(self, monkeypatch: pytest.MonkeyPatch) -> None:
        from app.modules.safety.business_agent.tools.hazard_action_tools import (
            send_hazard_urge,
        )

        _patch_urge_chain(
            monkeypatch, views=[_urge_view("recA")],
            urge_result={"total": 1, "recipients": 2, "sent": 1,
                         "skipped": 1, "errors": 0},
        )

        result = _run(send_hazard_urge(_ctx(None), record_id="recA"))

        assert result["sent"] == 1 and result["skipped"] == 1


# ── 确认流（挂起→确认/取消）——send_hazard_urge 作代表 ──────────────


class TestSendHazardUrgeConfirmationFlow:
    def _history_raw(self) -> list[dict]:
        from pydantic_ai.messages import (
            ModelMessagesTypeAdapter,
            ModelRequest,
            UserPromptPart,
        )

        base = json.loads(
            ModelMessagesTypeAdapter.dump_json(
                [ModelRequest(parts=[UserPromptPart(content="催一下这条隐患的责任人")])]
            )
        )
        from app.modules.safety.business_agent import executor
        from app.modules.safety.business_agent.schemas import PendingActionData

        pending = PendingActionData(
            tool_name="send_hazard_urge",
            arguments={"record_id": "recA"},
            summary="点对点催办未整改隐患 send_hazard_urge",
        )
        return executor._append_delegated_tool_call(base, pending)

    def _install_stubs(self, monkeypatch: pytest.MonkeyPatch) -> dict:
        from app.modules.safety.business_agent import executor

        captured: dict = {}

        class _FakeDb:
            async def execute(self, stmt: Any) -> Any:
                return SimpleNamespace(scalar_one=lambda: 0)

            async def flush(self) -> None:
                return None

            def add(self, obj: Any) -> None:
                return None

        async def fake_agent_run(user_message, deps, **kwargs):
            captured.update(kwargs)
            return SimpleNamespace(
                output="已发送催办卡", all_messages=lambda: [],
                new_messages=lambda: [],
            )

        async def fake_events(deps, result):
            return None

        async def fake_audit(*args, **kwargs):
            return None

        monkeypatch.setattr(executor, "_agent_run", fake_agent_run)
        monkeypatch.setattr(executor, "_record_tool_events", fake_events)
        monkeypatch.setattr(executor, "_audit_agent_resume", fake_audit)
        return captured

    def test_resume_approved_grants_tool(self, monkeypatch: pytest.MonkeyPatch) -> None:
        from app.modules.safety.business_agent import executor
        from app.modules.safety.business_agent.schemas import SafetyDeps

        captured = self._install_stubs(monkeypatch)
        sid = "11111111-2222-3333-4444-555555555555"
        deps = SafetyDeps(
            db=SimpleNamespace(), person=None, role="safety_admin",
            session_id=sid, channel="test",
        )

        out = _run(executor.resume(
            sid, deps, action_id="action-1", approved=True,
            message_history_raw=self._history_raw(),
        ))

        assert out["executed"] is True
        assert captured["resume_tool_names"] == ["send_hazard_urge"]
        results = captured["deferred_tool_results"]
        approved_ids = list(results.approvals.keys())
        assert len(approved_ids) == 1
        assert results.approvals[approved_ids[0]] is True

    def test_resume_cancelled_denies_tool(self, monkeypatch: pytest.MonkeyPatch) -> None:
        from pydantic_ai.tools import ToolDenied

        from app.modules.safety.business_agent import executor
        from app.modules.safety.business_agent.schemas import SafetyDeps

        captured = self._install_stubs(monkeypatch)
        sid = "11111111-2222-3333-4444-555555555555"
        deps = SafetyDeps(
            db=SimpleNamespace(), person=None, role="safety_admin",
            session_id=sid, channel="test",
        )

        _run(executor.resume(
            sid, deps, action_id="action-1", approved=False,
            message_history_raw=self._history_raw(),
        ))

        results = captured["deferred_tool_results"]
        value = list(results.approvals.values())[0]
        assert isinstance(value, ToolDenied)

    def test_resume_low_role_denied_with_reason(
        self, monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        """权限再校验：viewer 确认时被拒并说明（验收标准：角色不足被拒）。"""
        from pydantic_ai.tools import ToolDenied

        from app.modules.safety.business_agent import executor
        from app.modules.safety.business_agent.schemas import SafetyDeps

        captured = self._install_stubs(monkeypatch)
        sid = "11111111-2222-3333-4444-555555555555"
        deps = SafetyDeps(
            db=SimpleNamespace(), person=None, role="viewer",
            session_id=sid, channel="test",
        )

        _run(executor.resume(
            sid, deps, action_id="action-1", approved=True,
            message_history_raw=self._history_raw(),
        ))

        results = captured["deferred_tool_results"]
        value = list(results.approvals.values())[0]
        assert isinstance(value, ToolDenied)
        assert "权限不足" in str(value)


# ═══════════════════════════════════════════════════════════════════
# T3 通用推群 — push_report_to_chat
# ═══════════════════════════════════════════════════════════════════


class TestPushReportToChatWiring:
    def test_registry_and_groups(self) -> None:
        from app.modules.safety.business_agent.tools.hazard_action_tools import (
            push_report_to_chat,
        )

        assert registry.TOOL_KIND["push_report_to_chat"] is True
        assert registry._FUNC_BY_NAME["push_report_to_chat"] is push_report_to_chat
        assert "push_report_to_chat" not in registry.RESIDENT_TOOLS
        assert "push_report_to_chat" in tool_selector.TOOL_GROUPS["hazard"]

    def test_permissions_admin_and_leader_only(self) -> None:
        assert "push_report_to_chat" in allowed_tools(ROLE_SAFETY_ADMIN)
        assert "push_report_to_chat" in allowed_tools(ROLE_DEPT_LEADER)
        assert "push_report_to_chat" not in allowed_tools(ROLE_INSPECTOR)
        assert "push_report_to_chat" not in allowed_tools(ROLE_VIEWER)

    @pytest.mark.parametrize(
        ("message", "expected"),
        [
            ("把特殊作业日报发到群里", "push_report_to_chat"),
            ("把督办通报推送到群", "push_report_to_chat"),
            ("日报推送到群里", "push_report_to_chat"),
        ],
    )
    def test_route_selects_push_tool(self, message: str, expected: str) -> None:
        names = tool_selector.select_tool_names(message)
        assert expected in names, f"{message!r} → {names}"


class TestPushReportToChat:
    def test_supervision_routes_to_bulletin(
        self, monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        from app.modules.safety.business_agent.tools.hazard_action_tools import (
            push_report_to_chat,
        )

        calls: list[Any] = []

        async def _fake_send_bulletin(chat_id: str | None = None):
            calls.append(chat_id)
            return {"total": 3, "urgent": 1, "warning": 2, "marked": 0, "sent": True}

        monkeypatch.setattr(
            "app.modules.safety.service.hazard_direct.bulletin.send_bulletin",
            _fake_send_bulletin,
        )

        result = _run(push_report_to_chat(_ctx(None), report_kind="supervision",
                                          chat_id="oc_group1"))

        assert result["success"] is True
        assert calls == ["oc_group1"]
        assert result["total"] == 3

    def test_supervision_default_chat(self, monkeypatch: pytest.MonkeyPatch) -> None:
        from app.modules.safety.business_agent.tools.hazard_action_tools import (
            push_report_to_chat,
        )

        calls: list[Any] = []

        async def _fake_send_bulletin(chat_id: str | None = None):
            calls.append(chat_id)
            return {"total": 0, "urgent": 0, "warning": 0, "marked": 0, "sent": True}

        monkeypatch.setattr(
            "app.modules.safety.service.hazard_direct.bulletin.send_bulletin",
            _fake_send_bulletin,
        )

        _run(push_report_to_chat(_ctx(None), report_kind="supervision"))

        assert calls == [None]  # 缺省走系统配置群

    def test_daily_op_routes_with_target_chats(
        self, monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        from app.modules.safety.business_agent.tools.hazard_action_tools import (
            push_report_to_chat,
        )

        captured: dict[str, Any] = {}

        class _FakeService:
            def __init__(self, db: Any) -> None:
                pass

            async def generate_and_push(self, **kwargs: Any) -> Any:
                captured.update(kwargs)
                return SimpleNamespace(
                    report_date="2026-10-28", mode="today", total=5,
                    high_risk=1, medium_risk=2, low_risk=2, excluded=0,
                    push_results=[{"success": True}],
                    markdown_report="# 日报",
                )

        monkeypatch.setattr(
            "app.modules.safety.service.special_operation_daily_report"
            ".SpecialOperationDailyReportService",
            _FakeService,
        )

        class _FakeDb:
            async def commit(self) -> None:
                return None

        result = _run(push_report_to_chat(
            _ctx(_FakeDb()), report_kind="daily_op", chat_id="oc_group2",
        ))

        assert result["success"] is True
        assert captured["target_chats"] == ["oc_group2"]
        assert captured["push"] is True
        assert result["total"] == 5

    def test_alarm_fire_routes_to_fire_service(
        self, monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        from app.modules.safety.business_agent.tools.hazard_action_tools import (
            push_report_to_chat,
        )

        captured: dict[str, Any] = {}

        class _FakeFireService:
            def __init__(self, db: Any) -> None:
                pass

            async def generate_daily_report(self, **kwargs: Any) -> Any:
                captured.update(kwargs)
                return SimpleNamespace(
                    target_date="2026-10-28", total=4, analyzed=4,
                    push_results=[{"success": True}], markdown_report="# 报警日报",
                )

        monkeypatch.setattr(
            "app.modules.safety.service.fire_alarm.service.FireAlarmService",
            _FakeFireService,
        )

        class _FakeDb:
            async def commit(self) -> None:
                return None

        result = _run(push_report_to_chat(
            _ctx(_FakeDb()), report_kind="alarm", scope="fire",
            chat_id="oc_group3",
        ))

        assert result["success"] is True
        assert captured["chat_id"] == "oc_group3"
        assert captured["push"] is True
        assert result["total"] == 4

    def test_alarm_central_routes_to_central_service(
        self, monkeypatch: pytest.MonkeyPatch,
    ) -> None:
        from app.modules.safety.business_agent.tools.hazard_action_tools import (
            push_report_to_chat,
        )

        captured: dict[str, Any] = {}

        class _FakeCentralService:
            def __init__(self, db: Any) -> None:
                pass

            async def generate_daily_report(self, **kwargs: Any) -> Any:
                captured.update(kwargs)
                return SimpleNamespace(
                    target_date="2026-10-28", total=2, analyzed=2,
                    push_results=[{"success": True}], markdown_report="# 中控日报",
                )

        monkeypatch.setattr(
            "app.modules.safety.service.central_alarm.service.CentralAlarmService",
            _FakeCentralService,
        )

        class _FakeDb:
            async def commit(self) -> None:
                return None

        result = _run(push_report_to_chat(
            _ctx(_FakeDb()), report_kind="alarm", scope="central",
        ))

        assert result["success"] is True
        assert captured["chat_id"] is None  # 缺省走系统配置群

    def test_invalid_report_kind_hint(self) -> None:
        from app.modules.safety.business_agent.tools.hazard_action_tools import (
            push_report_to_chat,
        )

        result = _run(push_report_to_chat(_ctx(None), report_kind="moon"))

        assert result["success"] is False
        assert "report_kind" in result["error"]
        assert "supervision" in result["error"]

    def test_alarm_invalid_scope_hint(self) -> None:
        from app.modules.safety.business_agent.tools.hazard_action_tools import (
            push_report_to_chat,
        )

        result = _run(push_report_to_chat(_ctx(None), report_kind="alarm",
                                          scope="moon"))

        assert result["success"] is False
        assert "scope" in result["error"]
        assert "fire" in result["error"]


def _run(coro: Any) -> Any:
    import asyncio

    return asyncio.run(coro)


def _ctx(db: Any, *, chat_id: str | None = "oc_test") -> SimpleNamespace:
    return SimpleNamespace(
        deps=SimpleNamespace(db=db, chat_id=chat_id, role="safety_admin",
                             person=None),
    )


__all__ = [
    "TestSendHazardUrgeWiring",
    "TestSendHazardUrge",
    "TestSendHazardUrgeConfirmationFlow",
]
