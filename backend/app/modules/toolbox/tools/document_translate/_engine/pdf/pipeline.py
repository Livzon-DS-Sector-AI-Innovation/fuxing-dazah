"""PDF 翻译主流程：抽取 → 领域预读 → 复用共享执行器翻译 → 重排式 DOCX 输出 + 校验报告。

与 docx 流水线（cli.run_pipeline）共享：config/Settings、runner.translate_all
（并发/掩码/重试）、glossary、verify 报告机制。差异仅在"读"与"写"两端：
读端做结构化抽取（PDF 无段落对象），写端生成重排式新文档（PDF 无法原位回写）。
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from ..config import Settings
from ..glossary import Glossary, load_glossary
from ..runner import translate_all
from ..translate import LLMTranslator, MockTranslator, make_batches
from ..verify import VerifyContext, write_report
from .build import build_docx
from .extract import extract_pdf
from .profile import collect_profile_text, run_profile


def _parse_pages(spec: str) -> tuple[int, int] | None:
    """pages 配置 → 1-based 闭区间；空=全部页。格式错直接抛 ValueError。"""
    spec = (spec or "").strip()
    if not spec:
        return None
    m = re.fullmatch(r"(\d+)(?:-(\d+))?", spec)
    if not m:
        raise ValueError(f'pages 格式不正确 "{spec}"（应为 "1-20" 或 "6"）')
    a, b = int(m.group(1)), int(m.group(2) or m.group(1))
    return (min(a, b), max(a, b))


def run_pdf_pipeline(s: Settings, progress_cb=None) -> int:
    """执行 PDF 翻译完整流程，返回退出码（0=正常，2=存在必须处理的问题）。

    progress_cb(percent, message)：可选进度回调，各阶段/每批次上报（0~100）。
    """
    input_path = Path(s.input)
    suffix = "_bilingual" if s.mode == "bilingual" else "_translated"
    output_path = Path(s.output) if s.output else input_path.with_name(input_path.stem + suffix + ".docx")
    if output_path.resolve() == input_path.resolve():
        print("错误：输出路径不能与输入相同", file=sys.stderr)
        return 2
    report_path = Path(s.report) if s.report else output_path.with_suffix(".report.md")

    try:
        page_range = _parse_pages(s.pages)
    except ValueError as e:
        print(f"错误：{e}", file=sys.stderr)
        return 2

    glossary: Glossary = Glossary([])
    if s.glossary:
        try:
            glossary = load_glossary(s.glossary)
        except (ValueError, OSError) as e:
            print(f"错误：术语表加载失败：{e}", file=sys.stderr)
            return 2

    # ── 1. 结构化抽取 ─────────────────────────────────────────
    if progress_cb:
        progress_cb(3, "正在抽取 PDF 结构…")
    try:
        blocks, units, estats = extract_pdf(str(input_path), page_range)
    except ValueError as e:
        print(f"错误：{e}", file=sys.stderr)
        return 2

    todo, skipped_by_direction = [], 0
    for u in units:
        if s.direction == "zh2en" and u.lang != "zh":
            u.reason, skipped_by_direction = "非中文单元（zh2en 模式跳过）", skipped_by_direction + 1
            continue
        if s.direction == "en2zh" and u.lang != "en":
            u.reason, skipped_by_direction = "非英文单元（en2zh 模式跳过）", skipped_by_direction + 1
            continue
        u.action = "translate"
        todo.append(u)

    rng = f"第 {page_range[0]}~{page_range[1]} 页" if page_range else "全部页"
    print(f"结构抽取：共 {estats.total_pages} 页（本次处理 {rng}）｜标题 {estats.headings}｜"
          f"正文段落 {estats.paragraphs}｜表格 {estats.tables}（单元格 {estats.cells}）｜"
          f"目录页跳过 {estats.toc_pages}｜行号剥离 {estats.line_numbers_stripped}｜图片 {estats.images}（原位嵌入）")
    print(f"翻译单元：待翻译 {len(todo)} 个，方向不符跳过 {skipped_by_direction} 个")
    if progress_cb:
        progress_cb(8, f"结构抽取完成：共 {estats.total_pages} 页，待翻译 {len(todo)} 个单元")

    ctx = VerifyContext(
        input_path=str(input_path), output_path=str(output_path),
        mode=s.mode, direction=s.direction, doc_type=s.doc_type, mock=s.mock,
        source_kind="pdf",
        stats={
            "文档总页数": estats.total_pages,
            "本次处理范围": rng,
            "目录页（跳过，页码引用在重排版中无意义）": estats.toc_pages,
            "剥离行号（草案征询定位用）": estats.line_numbers_stripped,
            "剥离页眉/页脚的页数": estats.pages_with_header_stripped,
            "识别表格": estats.tables,
            "标题": estats.headings,
            "正文段落": estats.paragraphs,
            "表格单元格（非空）": estats.cells,
            "图片（不翻译，按原位置嵌入输出）": estats.images,
            "跨页断句合并": estats.cross_page_merges,
            "待翻译单元": len(todo),
            f"方向不符跳过（{s.direction}）": skipped_by_direction,
        },
    )

    # ── 2. 翻译器 ─────────────────────────────────────────────
    if s.mock:
        print("⚠️  MOCK 模式：不调用模型，译文=原文回显（用于流程自检）")
        translator = MockTranslator()
    else:
        if not s.api_key:
            print("错误：缺少 API Key。请在 config.json 中填写 api_key 字段", file=sys.stderr)
            return 2
        translator = LLMTranslator(s.api_base_url, s.api_key, s.model, max_retries=s.max_retries,
                                   thinking_level=s.thinking_level)

        # ── 2.5 领域预读：先读前 N 页生成领域档案，注入每批翻译提示 ──
        if s.profile_pages > 0:
            if progress_cb:
                progress_cb(10, f"领域预读：正在分析前 {s.profile_pages} 页生成领域档案…")
            first_page = page_range[0] if page_range else 1
            profile_src = collect_profile_text(blocks, first_page + s.profile_pages - 1)
            try:
                profile = run_profile(s.api_base_url, s.api_key, s.model, profile_src,
                                      thinking_level=s.thinking_level)
            except ValueError as e:
                print(f"错误：{e}", file=sys.stderr)
                return 2
            ctx.profile = profile
            print(f"领域档案（前 {s.profile_pages} 页自动分析）已生成，注入翻译提示。详见报告。")
            s.doc_type = s.doc_type + "\n\n【领域档案（基于文档开头自动分析）】\n" + profile

    # ── 3. 批量翻译（掩码 → 并发请求 → 还原校验，与 docx 流水线共用）──
    batches = make_batches(todo, max_units=s.batch_size)
    pos_of = {id(u): i for i, u in enumerate(todo)}

    def _translate_progress(frac: float, message: str) -> None:
        if progress_cb:
            progress_cb(15 + int(70 * frac), message)

    failed_units = translate_all(batches, translator, s, glossary, todo, pos_of, ctx,
                                 progress_cb=_translate_progress if progress_cb else None)

    translated_units = [u for u in todo if u.translation]
    ctx.stats["翻译成功"] = len(translated_units)
    ctx.stats["翻译失败/未回写"] = len(todo) - len(translated_units)

    # 术语一致性（提示级）
    if not s.mock and glossary:
        for u in translated_units:
            for problem in glossary.check(u.text, u.translation, u.lang):
                loc = "表格内" if u.in_table else "正文"
                ctx.add("术语不一致", f"#{u.idx}（{loc}）「{u.text[:36]}…」：{problem}")

    # ── 4. 生成重排式 DOCX ────────────────────────────────────
    if progress_cb:
        progress_cb(88, "正在生成重排式文档…")
    doc, expected_paras = build_docx(blocks, s.mode)
    try:
        doc.save(str(output_path))
    except OSError as e:
        print(f"错误：输出文件写入失败：{e}", file=sys.stderr)
        return 2

    # ── 5. 结构校验：重开文件核对段落数 + 每条译文确实落在输出中 ──
    if progress_cb:
        progress_cb(93, "正在结构校验…")
    structure_ok = _check_output(ctx, output_path, expected_paras, translated_units, s)
    if progress_cb:
        progress_cb(97, "正在生成校验报告…")
    write_report(ctx, str(report_path))

    hard = sum(1 for f in ctx.flags if f.category in {"无译文", "占位符异常", "翻译失败", "结构校验失败"})
    warn = len(ctx.flags) - hard
    print(f"\n输出：{output_path}")
    print(f"报告：{report_path}")
    print(f"结构校验：{'通过' if structure_ok else '失败'}｜硬性问题 {hard} 项｜提示 {warn} 项")
    if hard:
        print("存在必须人工处理的单元，详见报告。", file=sys.stderr)
    if progress_cb:
        progress_cb(100, "翻译完成")
    return 2 if (hard or not structure_ok) else 0


def _check_output(ctx: VerifyContext, output_path: Path, expected_paras: int,
                  translated_units, s: Settings) -> bool:
    """重开输出文件：段落数一致 + 译文完整性（空白不敏感包含检查）。"""
    from docx import Document
    from docx.oxml.ns import qn

    try:
        reopened = Document(str(output_path))
    except Exception as e:  # noqa: BLE001
        ctx.add("结构校验失败", f"输出文件无法重新解析: {e}")
        return False

    after = len(list(reopened.element.body.iter(qn("w:p"))))
    if after != expected_paras:
        ctx.add("结构校验失败", f"段落数不符：预期 {expected_paras}，实际 {after}")
        return False

    if s.mock:
        return True  # mock 译文=原文，包含检查无意义

    import re as _re

    all_text = _re.sub(r"\s+", "", "".join(reopened.element.body.itertext()))
    missing = [
        u for u in translated_units
        if _re.sub(r"\s+", "", u.translation)[:80] not in all_text
    ]
    for u in missing[:20]:
        loc = "表格内" if u.in_table else "正文"
        ctx.add("结构校验失败", f"#{u.idx}（{loc}）「{u.text[:36]}…」的译文未出现在输出文件中")
    if len(missing) > 20:
        ctx.add("结构校验失败", f"另有 {len(missing) - 20} 处译文缺失，从略")
    return not missing
