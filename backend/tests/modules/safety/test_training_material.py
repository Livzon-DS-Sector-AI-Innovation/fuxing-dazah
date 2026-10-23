"""链2 讲义/试卷生成：AI 插件 + 知识选择器 + docx 渲染 + 附件挂载链测试。

接缝（双接缝）：① opener 替身（training 客户端注入：list_all_records/
update_record/upload_media 三方法 fake）② ai_service 假服务（讲义+试卷
按序返回）。知识选择器走 db_session（regulation_chunks 真表，一次性测试
行用后即清）。规格 §5/§9。
"""

from __future__ import annotations

import io
from datetime import date
from typing import Any

import pytest
from docx import Document
from sqlalchemy import delete

from app.modules.safety.ai_training_material import (
    EXAM_QUESTION_KEYS,
    ExamInput,
    ExamPlugin,
    LectureInput,
    LecturePlugin,
    TrainingMaterialError,
    validate_exam,
)
from app.modules.safety.ai_training_material.docx_renderer import (
    render_exam_docx,
    render_lecture_docx,
)
from app.modules.safety.ai_training_material.prompts import (
    build_exam_user_prompt,
    resolve_exam_counts,
)
from app.modules.safety.ai_training_material.schemas import (
    LectureOutput,
    LectureSection,
)
from app.modules.safety.models.knowledge import RegulationChunk, SafetyKnowledgeArticle
from app.modules.safety.service.bitable_direct.errors import BitableConfigError
from app.modules.safety.service.training_behavior import material

_RAW_LECTURE = {
    "title": "受限空间作业安全讲义",
    "objectives": ["掌握受限空间作业审批流程", "能正确使用气体检测仪"],
    "sections": [
        {"heading": "一、作业前审批", "points": ["办理受限空间作业许可证", "气体检测合格后方可进入"]},
        {"heading": "二、作业中监护", "points": ["监护人不得离岗", "连续通风并定时复测"]},
    ],
    "key_takeaways": ["先审批后作业", "监护不断人"],
    "references": ["GB 30871-2022 §5.2"],
}

_RAW_EXAM = {
    "choice_questions": [
        {
            "question": "进入受限空间作业前必须办理什么手续？",
            "options": [
                {"label": "A", "text": "受限空间作业许可证"},
                {"label": "B", "text": "动火作业许可证"},
                {"label": "C", "text": "临时用电许可证"},
                {"label": "D", "text": "高处作业许可证"},
            ],
            "answer": "A",
            "source": "讲义要点",
        },
    ],
    "true_false_questions": [
        {"question": "受限空间作业监护人可以短暂离岗。", "answer": "错误", "source": "讲义要点"},
    ],
}


class FakeAIService:
    """单响应假服务（讲义或试卷单次调用场景）。"""

    def __init__(self, raw: dict | Exception = _RAW_LECTURE) -> None:
        self.raw = raw
        self.calls: list[dict] = []

    async def chat_parsed(self, *, messages, expected_keys, temperature=None):
        self.calls.append(
            {"messages": messages, "expected_keys": expected_keys,
             "temperature": temperature}
        )
        if isinstance(self.raw, Exception):
            raise self.raw
        return self.raw


class FakeSeqAIService:
    """按序多响应假服务（讲义→试卷连续调用场景）。"""

    def __init__(self, raws: list[dict | Exception]) -> None:
        self.raws = list(raws)
        self.calls: list[dict] = []

    async def chat_parsed(self, *, messages, expected_keys, temperature=None):
        self.calls.append(
            {"messages": messages, "expected_keys": expected_keys,
             "temperature": temperature}
        )
        raw = self.raws.pop(0) if self.raws else self.raws[-1]
        if isinstance(raw, Exception):
            raise raw
        return raw


class FakeTrainingClient:
    def __init__(self, records: list[dict] | None = None) -> None:
        self.records: list[dict] = records or []
        self.updated: list[tuple[str, dict]] = []
        self.uploads: list[tuple[str, str]] = []
        self.fail_upload = False
        self.upload_counter = 0

    async def list_all_records(self, **kwargs) -> list[dict]:
        return self.records

    async def update_record(self, record_id: str, fields: dict) -> None:
        self.updated.append((record_id, fields))

    async def upload_media(self, file_path: str, file_name: str) -> dict:
        if self.fail_upload:
            return {}
        self.upload_counter += 1
        self.uploads.append((file_path, file_name))
        return {"file_token": f"media-token-{self.upload_counter}", "name": file_name}


def _training_row(
    record_id: str = "recT1",
    *,
    name: str = "受限空间作业专项培训",
    remark: str | None = None,
    attachments: list[dict] | None = None,
) -> dict:
    fields: dict[str, Any] = {
        "培训名称": name,
        "培训类别": "专项培训",
        "培训形式": "线下",
        "组织部门": ["生产一部"],
        "培训学时": 4,
    }
    if remark is not None:
        fields["备注"] = remark
    if attachments is not None:
        fields["课件与试卷"] = attachments
    return {"record_id": record_id, "fields": fields}


def _opener(client: FakeTrainingClient):
    def open_client(domain: str, kind: str):
        if (domain, kind) != ("training_behavior", "training"):
            raise BitableConfigError(f"未配置: {domain}/{kind}")
        return client

    return open_client


_TEST_TITLE = "ZZ测试知识文档2099"


@pytest.fixture
async def knowledge_rows(db_session):
    """知识选择器测试夹具：一次性文章+条款行，用后即清（不留测试数据）。"""
    article = SafetyKnowledgeArticle(
        title=_TEST_TITLE, category="law", status="published",
    )
    db_session.add(article)
    await db_session.flush()
    chunks = [
        RegulationChunk(
            article_id=article.id, chunk_index=i,
            doc_title=_TEST_TITLE, doc_category="law",
            article_ref=f"第{i}条", chunk_text=f"测试条款{i}：进入受限空间前必须气体检测",
        )
        for i in range(1, 4)
    ]
    db_session.add_all(chunks)
    await db_session.flush()
    yield chunks
    await db_session.execute(
        delete(RegulationChunk).where(RegulationChunk.doc_title == _TEST_TITLE)
    )
    await db_session.execute(
        delete(SafetyKnowledgeArticle).where(SafetyKnowledgeArticle.title == _TEST_TITLE)
    )
    await db_session.commit()


# ── 讲义插件 ────────────────────────────────────────────────────


class TestLecturePlugin:
    async def test_generate_parses_lecture(self) -> None:
        service = FakeAIService()
        output = await LecturePlugin(service).generate(LectureInput(
            topic="受限空间作业专项培训",
            category="专项培训",
            hours="4",
            position="生产一部",
            knowledge_md="[1] GB 30871-2022\n    §5.2 受限空间作业应办理许可",
        ))
        assert output.title == "受限空间作业安全讲义"
        assert len(output.sections) == 2
        assert output.key_takeaways == ["先审批后作业", "监护不断人"]
        call = service.calls[0]
        assert call["temperature"] == 0.2
        assert call["expected_keys"] == [
            "title", "objectives", "sections", "key_takeaways", "references",
        ]
        user_msg = call["messages"][1]["content"]
        assert "主题：受限空间作业专项培训" in user_msg
        assert "GB 30871-2022" in user_msg

    async def test_missing_title_raises(self) -> None:
        raw = dict(_RAW_LECTURE, title="")
        with pytest.raises(TrainingMaterialError):
            await LecturePlugin(FakeAIService(raw)).generate(LectureInput())

    async def test_missing_sections_raises(self) -> None:
        raw = dict(_RAW_LECTURE, sections=[])
        with pytest.raises(TrainingMaterialError):
            await LecturePlugin(FakeAIService(raw)).generate(LectureInput())

    async def test_ai_failure_wraps(self) -> None:
        with pytest.raises(TrainingMaterialError):
            await LecturePlugin(FakeAIService(RuntimeError("boom"))).generate(
                LectureInput()
            )

    async def test_empty_knowledge_gets_placeholder(self) -> None:
        service = FakeAIService()
        await LecturePlugin(service).generate(LectureInput(topic="t"))
        assert "（无检索片段）" in service.calls[0]["messages"][1]["content"]


# ── docx 渲染 ───────────────────────────────────────────────────


class TestRenderDocx:
    def test_renders_key_sections(self) -> None:
        output = LectureOutput(
            title=_RAW_LECTURE["title"],
            objectives=_RAW_LECTURE["objectives"],
            sections=[
                LectureSection(heading=s["heading"], points=s["points"])
                for s in _RAW_LECTURE["sections"]
            ],
            key_takeaways=_RAW_LECTURE["key_takeaways"],
            references=_RAW_LECTURE["references"],
        )
        docx_bytes = render_lecture_docx(
            output, info_line="类别：专项培训｜学时：4｜对象：生产一部"
        )
        assert docx_bytes[:2] == b"PK"  # docx = zip
        doc = Document(io.BytesIO(docx_bytes))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "受限空间作业安全讲义" in text
        assert "培训目标" in text and "一、作业前审批" in text
        assert "核心要点" in text and "参考资料" in text
        assert "GB 30871-2022 §5.2" in text
        assert "类别：专项培训" in text


# ── 知识选择器（PG 真表） ───────────────────────────────────────


class TestKnowledgeSelector:
    async def test_match_and_format(self, db_session, knowledge_rows) -> None:
        markdown = await material.select_knowledge_fragments(
            db_session, topic="受限空间 ZZ测试知识文档2099 培训"
        )
        assert f"[1] {_TEST_TITLE}" in markdown
        assert "第1条" in markdown and "测试条款1" in markdown
        assert "\n---\n" in markdown or len(markdown) > 0

    async def test_no_match_returns_empty(self, db_session) -> None:
        markdown = await material.select_knowledge_fragments(
            db_session, topic="绝不匹配的主题ZZZZ"
        )
        assert markdown == ""

    async def test_truncation_cap(self, db_session, knowledge_rows) -> None:
        big = "长" * 9000
        knowledge_rows[0].chunk_text = big
        await db_session.flush()
        markdown = await material.select_knowledge_fragments(
            db_session, topic="ZZ测试知识文档2099"
        )
        assert 0 < len(markdown) <= 8000


# ── 生成编排（附件链 fake） ─────────────────────────────────────


class TestGenerateMaterialLecture:
    async def test_happy_path_mounts_attachment_and_marker(
        self, db_session
    ) -> None:
        client = FakeTrainingClient([_training_row(
            remark="既有备注",
            attachments=[{"file_token": "old-token", "name": "旧课件.docx"}],
        )])
        result = await material.generate_material(
            db_session, record_id="recT1", material_type="lecture",
            opener=_opener(client), ai_service=FakeAIService(),
        )
        today = date.today().isoformat()
        lecture = result["materials"][0]
        assert lecture["material_type"] == "lecture"
        # 上传形态：文件名=主题-讲义要点卡-日期.docx
        file_path, file_name = client.uploads[0]
        assert file_name == f"受限空间作业专项培训-讲义要点卡-{today}.docx"
        assert lecture["file_token"] == "media-token-1"
        # 写回形态：附件=既有+新增（不覆盖），备注=既有+追加标记
        record_id, fields = client.updated[0]
        assert record_id == "recT1"
        assert fields["课件与试卷"] == [
            {"file_token": "old-token", "name": "旧课件.docx"},
            {"file_token": "media-token-1", "name": file_name},
        ]
        assert fields["备注"] == f"既有备注\n{material.LECTURE_MARKER}({today})"

    async def test_idempotent_marker_blocks_regen(self, db_session) -> None:
        client = FakeTrainingClient([_training_row(
            remark=f"前次生成\n{material.LECTURE_MARKER}(2026-10-22)",
        )])
        with pytest.raises(material.MaterialStateError):
            await material.generate_material(
                db_session, record_id="recT1", material_type="lecture",
                opener=_opener(client), ai_service=FakeAIService(),
            )
        assert client.uploads == [] and client.updated == []

    async def test_row_not_found(self, db_session) -> None:
        client = FakeTrainingClient([_training_row("recOther")])
        with pytest.raises(material.MaterialStateError):
            await material.generate_material(
                db_session, record_id="recT1", material_type="lecture",
                opener=_opener(client), ai_service=FakeAIService(),
            )

    async def test_domain_not_connected(self, db_session) -> None:
        def opener(domain: str, kind: str):
            raise BitableConfigError("未接入")

        with pytest.raises(material.MaterialStateError):
            await material.generate_material(
                db_session, record_id="recT1", material_type="lecture",
                opener=opener, ai_service=FakeAIService(),
            )

    async def test_upload_failure_no_marker(self, db_session) -> None:
        client = FakeTrainingClient([_training_row()])
        client.fail_upload = True
        with pytest.raises(TrainingMaterialError):
            await material.generate_material(
                db_session, record_id="recT1", material_type="lecture",
                opener=_opener(client), ai_service=FakeAIService(),
            )
        assert client.updated == []  # 无标记写入 → 可重试

    async def test_ai_failure_no_writes(self, db_session) -> None:
        client = FakeTrainingClient([_training_row()])
        with pytest.raises(TrainingMaterialError):
            await material.generate_material(
                db_session, record_id="recT1", material_type="lecture",
                opener=_opener(client), ai_service=FakeAIService(RuntimeError("boom")),
            )
        assert client.uploads == [] and client.updated == []

    async def test_invalid_material_type(self, db_session) -> None:
        with pytest.raises(material.MaterialStateError):
            await material.generate_material(
                db_session, record_id="recT1", material_type="video",
            )


class TestGenerateMaterialExam:
    async def test_exam_material_and_marker(self, db_session) -> None:
        client = FakeTrainingClient([_training_row(remark="")])
        service = FakeSeqAIService([_RAW_LECTURE, _RAW_EXAM])
        result = await material.generate_material(
            db_session, record_id="recT1", material_type="exam",
            opener=_opener(client), ai_service=service,
        )
        today = date.today().isoformat()
        # 两次 AI：讲义（命题材料）→ 试卷
        assert len(service.calls) == 2
        exam_user_msg = service.calls[1]["messages"][1]["content"]
        assert "材料：" in exam_user_msg and "讲义要点" in exam_user_msg
        assert "检索片段" in exam_user_msg
        # 试卷 temp 0.1 + 默认 expected_keys 只含 单选/判断（多选/填空=0 不列）
        assert service.calls[1]["temperature"] == 0.1
        assert service.calls[1]["expected_keys"] == [
            "choice_questions", "true_false_questions",
        ]
        exam = result["materials"][0]
        assert exam["material_type"] == "exam"
        assert exam["file_name"] == f"受限空间作业专项培训-试卷-{today}.docx"
        _, fields = client.updated[0]
        assert fields["备注"] == f"{material.EXAM_MARKER}({today})"

    async def test_both_generates_two_files_one_writeback(
        self, db_session
    ) -> None:
        client = FakeTrainingClient([_training_row(remark="")])
        service = FakeSeqAIService([_RAW_LECTURE, _RAW_EXAM])
        result = await material.generate_material(
            db_session, record_id="recT1", material_type="both",
            opener=_opener(client), ai_service=service,
        )
        today = date.today().isoformat()
        assert len(client.uploads) == 2  # 讲义+试卷两文件
        assert len(client.updated) == 1  # 一次写回
        _, fields = client.updated[0]
        assert fields["备注"] == (
            f"{material.LECTURE_MARKER}({today})\n{material.EXAM_MARKER}({today})"
        )
        assert {m["material_type"] for m in result["materials"]} == {"lecture", "exam"}
        # 讲义 AI 只跑一次（讲义 docx 与试卷命题共用）
        assert len(service.calls) == 2

    async def test_exam_idempotent_block(self, db_session) -> None:
        client = FakeTrainingClient([_training_row(
            remark=f"{material.EXAM_MARKER}(2026-10-22)",
        )])
        with pytest.raises(material.MaterialStateError):
            await material.generate_material(
                db_session, record_id="recT1", material_type="exam",
                opener=_opener(client),
                ai_service=FakeSeqAIService([_RAW_LECTURE, _RAW_EXAM]),
            )


class TestListMaterials:
    async def test_lists_attachments_and_generations(self) -> None:
        today = date.today().isoformat()
        client = FakeTrainingClient([_training_row(
            remark=f"备注\n{material.LECTURE_MARKER}({today})",
            attachments=[
                {"file_token": "t1", "name": "受限空间-讲义要点卡-2026.docx"},
                {"file_token": "t2", "name": "受限空间-试卷-2026.docx"},
                {"file_token": "t3", "name": "其他附件.pdf"},
            ],
        )])
        data = await material.list_materials(
            record_id="recT1", opener=_opener(client)
        )
        types = {a["name"]: a["material_type"] for a in data["attachments"]}
        assert types["受限空间-讲义要点卡-2026.docx"] == "lecture"
        assert types["受限空间-试卷-2026.docx"] == "exam"
        assert types["其他附件.pdf"] is None
        assert data["generations"] == [{"material_type": "lecture", "generated_at": today}]


# ── 试卷提示词拼装（§5.5 句式） ─────────────────────────────────


class TestExamPrompt:
    def test_default_config_assembly(self) -> None:
        prompt = build_exam_user_prompt(material="材料内容", config=None)
        assert "出 5 道单选题（4 选项 A/B/C/D，只有一个正确答案）" in prompt
        assert "出 5 道判断题（answer 只能是\"正确\"或\"错误\"，正确与错误大致各半）" in prompt
        assert "多选题" not in prompt and "填空题" not in prompt
        assert '"choice_questions"' in prompt
        assert '"multi_choice_questions"' not in prompt
        assert prompt.startswith("请根据以下材料生成试卷：出 5 道单选题")
        assert "输出格式：{ " in prompt and " }" in prompt
        assert prompt.endswith("材料：\n材料内容")

    def test_full_config_assembly(self) -> None:
        prompt = build_exam_user_prompt(
            material="m",
            config={"choice_count": 4, "true_false_count": 2,
                    "multi_choice_count": 3, "fill_blank_count": 1},
        )
        assert "出 4 道单选题" in prompt
        assert "出 2 道判断题" in prompt
        assert "出 3 道多选题（4 选项，有 2-3 个正确答案）" in prompt
        assert "出 1 道填空题（将关键知识点挖空，question 中用______表示空位，question 里不能出现答案）" in prompt
        for key in EXAM_QUESTION_KEYS:
            assert f'"{key}"' in prompt

    def test_fill_count_alias_accepted(self) -> None:
        """规格 §5.5 行文键 fill_count 与 HR 实键 fill_blank_count 等价。"""
        by_alias = resolve_exam_counts({"fill_count": 2})
        by_canonical = resolve_exam_counts({"fill_blank_count": 2})
        assert by_alias["fill_blank_questions"] == 2
        assert by_canonical["fill_blank_questions"] == 2

    def test_zero_all_config_raises(self) -> None:
        with pytest.raises(ValueError):
            build_exam_user_prompt(
                material="m",
                config={"choice_count": 0, "true_false_count": 0},
            )

    def test_resolve_counts_defaults_and_invalid(self) -> None:
        assert resolve_exam_counts(None) == {
            "choice_questions": 5, "true_false_questions": 5,
            "multi_choice_questions": 0, "fill_blank_questions": 0,
        }
        counts = resolve_exam_counts({"choice_count": -3, "true_false_count": "x"})
        assert counts["choice_questions"] == 0
        assert counts["true_false_questions"] == 5


# ── 试卷校验（对齐 HR _validate_exam 语义） ─────────────────────


class TestValidateExam:
    def test_valid_structure_passes_with_source(self) -> None:
        validated = validate_exam(_RAW_EXAM)
        assert validated["choice_questions"][0]["answer"] == "A"
        assert validated["choice_questions"][0]["source"] == "讲义要点"
        assert validated["true_false_questions"][0]["answer"] == "错误"
        assert validated["multi_choice_questions"] == []

    def test_true_false_answer_normalized(self) -> None:
        raw = {"true_false_questions": [
            {"question": "q1", "answer": "对"},
            {"question": "q2", "answer": "×"},
        ]}
        validated = validate_exam(raw, {"choice_count": 0, "true_false_count": 2})
        assert [q["answer"] for q in validated["true_false_questions"]] == ["正确", "错误"]

    def test_true_false_illegal_answer_raises(self) -> None:
        raw = {"true_false_questions": [{"question": "q", "answer": "也许"}]}
        with pytest.raises(ValueError, match="判断题答案不合法"):
            validate_exam(raw, {"choice_count": 0, "true_false_count": 1})

    def test_multi_answer_normalized_and_validated(self) -> None:
        raw = {"multi_choice_questions": [{
            "question": "q",
            "options": [
                {"label": "A", "text": "a"}, {"label": "B", "text": "b"},
                {"label": "C", "text": "c"}, {"label": "D", "text": "d"},
            ],
            "answer": "a，b",
        }]}
        validated = validate_exam(raw, {"choice_count": 0, "true_false_count": 0,
                                        "multi_choice_count": 1})
        assert validated["multi_choice_questions"][0]["answer"] == "A,B"

    def test_answer_not_in_labels_raises(self) -> None:
        raw = {"choice_questions": [{
            "question": "q",
            "options": [{"label": "A", "text": "a"}, {"label": "B", "text": "b"}],
            "answer": "C",
        }]}
        with pytest.raises(ValueError, match="答案与选项不匹配"):
            validate_exam(raw, {"choice_count": 1, "true_false_count": 0})

    def test_missing_question_raises(self) -> None:
        raw = {"choice_questions": [{"question": "  ", "answer": "A"}]}
        with pytest.raises(ValueError, match="缺少题干"):
            validate_exam(raw, {"choice_count": 1, "true_false_count": 0})

    def test_empty_requested_key_raises(self) -> None:
        with pytest.raises(ValueError, match="未生成任何"):
            validate_exam({"choice_questions": []}, None)

    def test_excess_truncated(self) -> None:
        raw = {"true_false_questions": [
            {"question": f"q{i}", "answer": "正确"} for i in range(5)
        ]}
        validated = validate_exam(
            raw, {"choice_count": 0, "true_false_count": 2}
        )
        assert len(validated["true_false_questions"]) == 2


class TestExamPlugin:
    async def test_expected_keys_exclude_zero_counts(self) -> None:
        service = FakeAIService(_RAW_EXAM)
        output = await ExamPlugin(service).generate(ExamInput(
            material="材料",
            config={"choice_count": 1, "true_false_count": 1,
                    "multi_choice_count": 0, "fill_count": 0},
        ))
        call = service.calls[0]
        assert call["temperature"] == 0.1
        assert call["expected_keys"] == ["choice_questions", "true_false_questions"]
        assert output.total_count == 2

    async def test_validation_failure_wraps(self) -> None:
        bad = {"choice_questions": [{"question": "q", "answer": "Z"}],
               "true_false_questions": [{"question": "t", "answer": "正确"}]}
        with pytest.raises(TrainingMaterialError, match="校验失败"):
            await ExamPlugin(FakeAIService(bad)).generate(ExamInput(material="m"))

    async def test_ai_failure_wraps(self) -> None:
        with pytest.raises(TrainingMaterialError):
            await ExamPlugin(FakeAIService(RuntimeError("boom"))).generate(
                ExamInput(material="m")
            )


class TestRenderExamDocx:
    def test_renders_sections_and_answers(self) -> None:
        questions = validate_exam(_RAW_EXAM, {
            "choice_count": 1, "true_false_count": 1,
            "multi_choice_count": 0, "fill_blank_count": 0,
        })
        questions["fill_blank_questions"] = [{
            "question": "进入受限空间前必须进行______检测", "answer": "气体",
        }]
        docx_bytes = render_exam_docx(questions, title="受限空间作业考核试卷")
        doc = Document(io.BytesIO(docx_bytes))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert "受限空间作业考核试卷" in text
        assert "一、单选题（共1题）" in text
        assert "二、判断题（共1题）" in text
        assert "四、填空题（共1题）" in text
        assert "三、多选题" not in text  # 空节不渲染
        assert "A. 受限空间作业许可证" in text
        assert "对 □    错 □" in text
        assert "参考答案" in text and "一、单选题答案：" in text
        assert "进入受限空间前必须进行______检测" in text
