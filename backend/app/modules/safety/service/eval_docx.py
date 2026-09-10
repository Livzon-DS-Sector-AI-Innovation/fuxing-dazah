"""演练评估表 docx 渲染 — 纯函数。

把 DrillEvalOutput 渲染为与人工样例同格式的评估表 docx（bytes）：
标题 + 单张表格（基本信息区 / 演练类型勾选 / 总体评价 / 演练等级评定勾选 /
问题整改明细 / 评估人落款）。无 DB / 网络依赖，便于直测。
"""

from __future__ import annotations

import io

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from app.modules.safety.ai_drill_eval.schemas import DrillEvalOutput

_NCOLS = 7
_EMPTY = "—"


def map_drill_type(hint: str) -> str:
    """把主表「演练类型」文本映射到四选一枚举（无法映射返回空串）。"""
    hint = (hint or "").strip()
    if not hint:
        return ""
    if "现场" in hint:
        return "现场处置演练"
    if "桌面" in hint:
        return "桌面演练"
    if "综合" in hint:
        return "综合演练"
    if "专项" in hint:
        return "专项演练"
    return ""


def _checked_line(options: list[str], checked: str) -> str:
    """生成勾选行文本，如「□综合演练 □专项演练 ☑现场处置演练 □桌面演练」。

    checked 支持前缀匹配（等级选项带分数后缀：「良好」命中「良好（80-89分）」）。
    """
    return "  ".join(
        ("☑" if checked and opt.startswith(checked) else "□") + opt for opt in options
    )


def _set_cell_text(cell, text: str, *, bold: bool = False, size: int = 9) -> None:
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold else WD_ALIGN_PARAGRAPH.LEFT
    para.paragraph_format.space_before = Pt(1)
    para.paragraph_format.space_after = Pt(1)
    run = para.add_run(text)
    run.font.bold = bold
    run.font.size = Pt(size)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def _merge_row(row, spans: list[tuple[int, int, str]], *, bold: bool = False) -> None:
    """按 (起始列, 结束列含, 文本) 合并并填文本。"""
    for start, end, text in spans:
        cell = row.cells[start]
        if end > start:
            cell = cell.merge(row.cells[end])
        _set_cell_text(cell, text, bold=bold)


def render_eval_docx(
    output: DrillEvalOutput,
    *,
    doc_title: str,
    drill_type_hint: str = "",
) -> bytes:
    """渲染评估表 docx。

    Args:
        output: 评估表结构化内容
        doc_title: 完整标题（编排层拼装，如「提炼工程二部妥布氨水泄漏现场处置演练评估表」）
        drill_type_hint: 主表「演练类型」原文，output.drill_type_check 为空时用于映射勾选
    """
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)

    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run(doc_title)
    title_run.font.bold = True
    title_run.font.size = Pt(14)
    title_run.font.name = "宋体"
    title_run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    type_check = output.drill_type_check or map_drill_type(drill_type_hint)
    grade = output.grade

    n_issues = max(len(output.issues), 1)
    table = doc.add_table(rows=4 + 1 + 2 + 2 + 1 + n_issues + 1, cols=_NCOLS)
    table.style = "Table Grid"

    # ── 基本信息区（4 行：标签|值×2|标签|值×3）──
    info_rows: list[tuple[str, str, str, str]] = [
        ("演练名称", output.drill_name or _EMPTY, "演练时间", output.drill_time or _EMPTY),
        ("演练地点", output.drill_place or _EMPTY, "组织部门", output.org_department or _EMPTY),
        ("现场指挥员", output.commander or _EMPTY, "记录人", output.recorder or _EMPTY),
        ("参加人员", output.participants_desc or _EMPTY, "参演人数", output.participant_count_desc or _EMPTY),
    ]
    for r, (label1, value1, label2, value2) in enumerate(info_rows):
        row = table.rows[r]
        _merge_row(row, [(0, 0, label1), (1, 2, value1), (3, 3, label2), (4, 6, value2)], bold=True)

    # ── 演练类型勾选行 ──
    r = 4
    _merge_row(
        table.rows[r],
        [(0, 6, _checked_line(["综合演练", "专项演练", "现场处置演练", "桌面演练"], type_check))],
    )

    # ── 总体评价（标签行 + 内容行）──
    _merge_row(table.rows[5], [(0, 6, "总体评价")], bold=True)
    comment_row = table.rows[6]
    cell = comment_row.cells[0].merge(comment_row.cells[6])
    paragraphs = [p for p in output.overall_comment.split("\n") if p.strip()] or [_EMPTY]
    _set_cell_text(cell, paragraphs[0])
    for extra in paragraphs[1:]:
        para = cell.add_paragraph()
        run = para.add_run(extra)
        run.font.size = Pt(9)
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    # ── 演练等级评定（标签行 + 勾选行）──
    _merge_row(table.rows[7], [(0, 6, "演练等级评定")], bold=True)
    _merge_row(
        table.rows[8],
        [(0, 6, _checked_line(["优秀（90-100分）", "良好（80-89分）", "合格（60-79分）", "不合格（60分以下）"], grade))],
    )

    # ── 问题整改明细 ──
    header_row = table.rows[9]
    for c, text in enumerate(["序号", "存在问题", "整改措施", "完成期限", "完成情况", "整改人", "确认人"]):
        _set_cell_text(header_row.cells[c], text, bold=True)

    for i, item in enumerate(output.issues, start=1):
        row = table.rows[9 + i]
        values = [
            str(i),
            item.issue or _EMPTY,
            item.action or _EMPTY,
            item.deadline or _EMPTY,
            item.done_status or _EMPTY,
            item.rectifier or _EMPTY,
            item.confirmer or _EMPTY,
        ]
        for c, text in enumerate(values):
            _set_cell_text(row.cells[c], text)

    # ── 落款行：评估人 | 姓名 | 评估日期 | 日期 ──
    sign_row = table.rows[10 + n_issues]
    _merge_row(
        sign_row,
        [
            (0, 0, "评估人"),
            (1, 2, output.evaluator or _EMPTY),
            (3, 3, "评估日期"),
            (4, 6, output.eval_date or _EMPTY),
        ],
        bold=True,
    )

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
