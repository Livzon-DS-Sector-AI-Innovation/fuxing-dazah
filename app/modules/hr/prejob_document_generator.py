"""Generate pre-job training plan documents from templates."""

from io import BytesIO
from pathlib import Path

import openpyxl

from app.modules.hr.models import Employee


def _find_template() -> Path:
    """Locate the xlsx template, trying several path candidates."""
    candidates = [
        Path("assets/hr/7.4岗前培训计划.xlsx"),
        Path("../assets/hr/7.4岗前培训计划.xlsx"),
        Path(__file__).resolve().parent.parent.parent.parent
        / "assets/hr"
        / "7.4岗前培训计划.xlsx",
    ]
    for p in candidates:
        if p.exists():
            return p
    raise FileNotFoundError("模板文件未找到: 7.4岗前培训计划.xlsx")


DEPT_CONTENT_MAP: dict[str, list[str]] = {
    "人事行政部": [
        "公司级公用文件(详见附件一)",
        "部门级公用文件(详见附件二)",
        "人事行政部人事行政专员岗位文件(详见附件三)",
        "人事行政专员岗位职责(QP.PM.053)",
        "生产安全知识",
        "岗前培训计划",
    ],
}


def generate_prejob_training_plan(employee: Employee) -> BytesIO:
    """Fill the pre-job training plan template with employee data.

    Returns a BytesIO buffer containing the generated xlsx.
    """
    template_path = _find_template()
    wb = openpyxl.load_workbook(str(template_path))
    ws = wb.active

    # Part 1: Employee overview
    ws["C5"] = employee.name or ""
    ws["I5"] = employee.department or ""
    ws["C6"] = employee.employee_number or ""
    ws["I6"] = str(employee.hire_date) if employee.hire_date else ""
    ws["C7"] = employee.position or ""

    # Part 2: Training content (auto-fill by department)
    content_list = DEPT_CONTENT_MAP.get(employee.department or "", [])
    for i, content in enumerate(content_list):
        row = 11 + i
        if row <= 20:
            ws[f"B{row}"] = content

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
