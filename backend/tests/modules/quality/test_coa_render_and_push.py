"""COA 渲染链路与推送幂等测试（此前 _render_cao_file/_mark_pushed 零覆盖）。"""

import pytest

# ── _render_cao_file：模板填充落盘 ──


def test_render_coa_file_fills_placeholders():
    from app.modules.quality.api._common import _render_cao_file
    from app.modules.quality.storage import REPORT_TEMPLATE_DIR

    fill_data = {
        "批号": "HAFTEST001",
        "产品名称": "测试产品",
        "流水号": "26092301",
        "水分": 2.5,
        "水分_判定": "合格",
    }
    output_path, output_filename, content = _render_cao_file(
        REPORT_TEMPLATE_DIR / "3205.docx",
        fill_data,
        "测试产品",
        "HAFTEST001",
    )
    try:
        assert output_path.exists()
        assert output_filename.startswith("COA-测试产品-HAFTEST001")
        assert content  # 文件非空
        # 流水号应已填充进文档（docx 内部是 zip，直接读段落验证）
        from docx import Document

        doc = Document(str(output_path))
        texts = "\n".join(p.text for p in doc.paragraphs)
        assert "26092301" in texts
    finally:
        output_path.unlink(missing_ok=True)


def test_render_coa_file_unique_filename_with_serial():
    from app.modules.quality.api._common import _render_cao_file
    from app.modules.quality.storage import REPORT_TEMPLATE_DIR

    p1, f1, _ = _render_cao_file(
        REPORT_TEMPLATE_DIR / "3205.docx",
        {"批号": "HAFTEST002", "流水号": "26092302"},
        "测试产品", "HAFTEST002",
    )
    p2, f2, _ = _render_cao_file(
        REPORT_TEMPLATE_DIR / "3205.docx",
        {"批号": "HAFTEST002", "流水号": "26092303"},
        "测试产品", "HAFTEST002",
    )
    try:
        # 同批号不同流水号 → 文件名不同（防同名覆盖错单）
        assert f1 != f2
    finally:
        p1.unlink(missing_ok=True)
        p2.unlink(missing_ok=True)


# ── 推送幂等标记 ──


@pytest.mark.anyio
async def test_mark_pushed_blocks_second_call():
    from app.modules.quality.feishu.daily_push import _mark_pushed

    key = "test:mark_pushed"
    try:
        from redis.exceptions import RedisError
    except ImportError:
        pytest.skip("redis 未安装")
    try:
        from app.core.redis import cache_delete

        # acquire_lock 的键带 "lock:" 前缀，清理时也要带
        await cache_delete(f"lock:quality:push:{key}")
        assert await _mark_pushed(key) is True
        assert await _mark_pushed(key) is False  # 第二次调用跳过（幂等）
    except RedisError:
        pytest.skip("Redis 不可用，跳过幂等标记测试")


# ── 存储降级：MinIO 未启用时纯本地 ──


def test_upload_attachment_local_fallback(monkeypatch, tmp_path):
    from app.modules.quality import storage as qs

    monkeypatch.setattr(qs, "ATTACH_LOCAL_DIR", tmp_path)
    qs.upload_attachment("t1/file.txt", b"hello", "text/plain")
    assert (tmp_path / "t1/file.txt").read_bytes() == b"hello"
    assert qs.read_attachment("t1/file.txt") == b"hello"
    qs.delete_attachment("t1/file.txt")
    assert qs.read_attachment("t1/file.txt") is None
