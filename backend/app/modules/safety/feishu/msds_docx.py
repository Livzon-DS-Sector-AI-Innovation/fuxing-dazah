"""MSDS 标准模板 docx 生成器 — 填充企业标准模板。

将 AI 提取的 28 字段填充到《MSDS标准模板.docx》（单表 11×8 合并单元格布局：
名称/UN No/分子式/CAS No + 危险性说明标签要素 + 理化特性/职业接触限值交错两栏
+ 健康危害/急救/接触控制 + 稳定性/泄漏/消防/废弃 + 操作储存）。

填充策略：
  - 直接字段（名称/理化特性/限值/健康危害等）→ 对应 label 后的【】填值
  - 组合字段（急救/接触控制/稳定性/消防/废弃/操作储存）→ 按子标签拆分逐【】填充；
    子标签无法匹配时留空，单占位符节（泄漏应急处理）整块填充
  - 标签要素模板无占位符 → 仅平台存档，docx 不输出
"""

from __future__ import annotations

import logging
import re
from pathlib import Path

from docx import Document

logger = logging.getLogger(__name__)

_TEMPLATE_PATH = Path(__file__).resolve().parents[1] / "templates" / "msds_standard_template.docx"


# ── 直接字段：模板 label（归一化）→ entry 字段 ──

_DIRECT_LABEL_TO_FIELD: dict[str, str] = {
    "名称": "name",
    "UNNo": "un_no",
    "分子式": "molecular_formula",
    "CASNo": "cas_no",
    "危险性说明": "hazard_statement",
    "外观与现状": "appearance",
    "溶解性": "solubility",
    "熔点": "melting_point",
    "闪点": "flash_point",
    "沸点": "boiling_point",
    "相对密度（水）": "relative_density",
    "爆炸上限(%)": "explosion_upper_limit",
    "爆炸下限(%)": "explosion_lower_limit",
    "自燃温度": "autoignition_temperature",
    "分解温度": "decomposition_temperature",
    "时间加权平均容许浓度（PC-TWA）": "pc_twa",
    "短时间接触容许浓度（PC-STEL）": "pc_stel",
    "最高容许浓度（MAC）": "mac",
    "健康危害": "health_hazard",
}

# ── 组合节：模板子标签 → 所属节 + AI 输出中的别名 ──
# aliases 为在 AI 组合文本中用于拆分/匹配的子标签；空 aliases 表示整块填充

_SUB_LABEL_SPECS: dict[str, tuple[str, list[str]]] = {
    # 急救措施（first_aid）
    "皮肤接触": ("急救措施", ["皮肤接触"]),
    "眼睛接触": ("急救措施", ["眼睛接触"]),
    "吸入": ("急救措施", ["吸入"]),
    "食入": ("急救措施", ["食入"]),
    "对医生的特别提示": ("急救措施", ["对医生的特别提示"]),
    # 接触控制/个体防护（exposure_controls）
    "监测方法": ("接触控制/个体防护", ["监测方法"]),
    "工程控制": ("接触控制/个体防护", ["工程控制"]),
    "呼吸系统防护": ("接触控制/个体防护", ["呼吸系统防护"]),
    "眼睛防护": ("接触控制/个体防护", ["眼睛防护"]),
    "身体防护": ("接触控制/个体防护", ["皮肤和身体防护", "身体防护"]),
    "手防护": ("接触控制/个体防护", ["手防护"]),
    "其他防护": ("接触控制/个体防护", ["其他防护"]),
    # 稳定性和反应性（stability_reactivity）
    "稳定性": ("稳定性和反应性", ["稳定性"]),
    "禁配物": ("稳定性和反应性", ["禁配物"]),
    "聚合危害": ("稳定性和反应性", ["聚合危害"]),
    # 泄漏应急处理（leakage_response）— 整块填充
    "应急处理": ("泄漏应急处理", []),
    # 消防措施（fire_fighting）
    "危险特性": ("消防措施", ["特别危险性", "危险特性"]),
    "有害燃烧产物": ("消防措施", ["有害燃烧产物"]),
    "灭火方法及灭火剂": ("消防措施", ["灭火方法及灭火剂", "灭火方法", "灭火剂"]),
    "灭火注意事项": ("消防措施", ["灭火注意事项及防护措施", "灭火注意事项"]),
    # 废弃处置（waste_disposal）
    "废弃物性质": ("废弃处置", ["废弃化学品", "废弃物性质"]),
    "废弃方法": ("废弃处置", ["废弃方法", "废弃注意事项"]),
    "包装": ("废弃处置", ["污染包装物", "包装"]),
    # 操作处置与储存注意事项（handling_storage）
    "操作注意事项": ("操作处置与储存注意事项", ["操作注意事项"]),
    "储存注意事项": ("操作处置与储存注意事项", ["储存注意事项"]),
}

# 节 → 组合字段 key
_SECTION_FIELD: dict[str, str] = {
    "急救措施": "first_aid",
    "接触控制/个体防护": "exposure_controls",
    "稳定性和反应性": "stability_reactivity",
    "泄漏应急处理": "leakage_response",
    "消防措施": "fire_fighting",
    "废弃处置": "waste_disposal",
    "操作处置与储存注意事项": "handling_storage",
}

# 节标题（无占位符，用于 last_label 上下文）
_SECTION_HEADINGS: set[str] = set(_SECTION_FIELD.keys()) | {"急救措施"}


# ── 文本工具 ──


def _normalize(text: str) -> str:
    """归一化标签：去空白 + 去全/半角冒号。"""
    t = re.sub(r"\s+", "", text)
    t = t.rstrip("：:")
    return t


def _label_before_placeholder(text: str) -> str:
    """取占位符【】之前的标签部分（归一化前）。"""
    idx = text.find("【")
    return text[:idx] if idx >= 0 else ""


def _split_into_parts(text: str, aliases: list[str]) -> dict[str, str]:
    """按子标签+冒号边界拆分组合文本 → {alias: 子文本}。

    要求子标签后必须有冒号（全/半角）才视为分隔符，避免把句子中
    嵌入的标签词（如"与强氧化剂等禁配物接触"里的"禁配物"）误切。
    """
    if not text:
        return {}
    if not aliases:
        return {}
    joined = "|".join(re.escape(a) for a in sorted(aliases, key=len, reverse=True))
    sep_re = re.compile(f"({joined})[：:]")
    tokens = sep_re.split(text)
    result: dict[str, str] = {}
    current: str | None = None
    for tok in tokens:
        t = tok.strip()
        m = re.fullmatch(f"({joined})[：:]?", t)
        if m:
            current = m.group(1).strip("：:").strip()
            result.setdefault(current, "")
        elif current is not None and t:
            result[current] = (result[current] + t).strip()
    return {k: v for k, v in result.items() if v}


def _replace_placeholder(paragraph, value: str) -> bool:
    """替换段落中第一个【...】为 value；值置入首个 run 保持格式。"""
    full = "".join(r.text for r in paragraph.runs)
    if "【" not in full:
        return False
    value = (value or "").strip()
    # 模板占位符后自带单位" mg/m3"，若值本身已含单位则去掉模板后缀，避免"350mg/m3 mg/m3"
    if re.search(r"【[^】]*】\s*mg/m3\s*$", full) and re.search(r"mg/m3\s*$", value, re.IGNORECASE):
        full = re.sub(r"\s*mg/m3\s*$", "", full)
    new_full = re.sub(r"【[^】]*】", value, full, count=1)
    if paragraph.runs:
        paragraph.runs[0].text = new_full
        for r in paragraph.runs[1:]:
            r.text = ""
    return True


# ── 值解析 ──


def _resolve_value(
    label: str,
    last_label: str,
    entry: dict,
    section_parts: dict[str, dict[str, str]],
) -> str:
    """解析占位符应填的值。"""
    # ① 直接字段
    if label in _DIRECT_LABEL_TO_FIELD:
        val = entry.get(_DIRECT_LABEL_TO_FIELD[label]) or ""
        if label == "相对密度（水）":
            # 模板 label 已含"（水）"，AI 值常含"（水=1）"，去重避免"相对密度（水）：0.786（水=1）"
            val = val.replace("（水=1）", "").replace("(水=1)", "")
        return val
    # ② 组合子标签
    if label in _SUB_LABEL_SPECS:
        section, aliases = _SUB_LABEL_SPECS[label]
        if not aliases:
            # 单占位符节（应急处理）→ 整块填充
            return entry.get(_SECTION_FIELD.get(section, "")) or ""
        parts = section_parts.get(section, {})
        for a in aliases:
            if a in parts:
                return parts[a]
        return ""
    # ③ 无本段标签 → 用上一个标签（名称/UN No/危险性说明/健康危害的独立【】段）
    if last_label in _DIRECT_LABEL_TO_FIELD:
        return entry.get(_DIRECT_LABEL_TO_FIELD[last_label]) or ""
    if last_label in _SUB_LABEL_SPECS:
        section, aliases = _SUB_LABEL_SPECS[last_label]
        if not aliases:
            return entry.get(_SECTION_FIELD.get(section, "")) or ""
        parts = section_parts.get(section, {})
        for a in aliases:
            if a in parts:
                return parts[a]
    return ""


def build_standard_msds_docx(entry: dict, output_path: str, template_path: str | None = None) -> str:
    """将单条 28 字段填充到企业标准模板，输出 docx。

    Args:
        entry: AI 提取的单化学品 28 字段 dict
        output_path: 输出 docx 路径
        template_path: 模板路径（默认使用安全模块内置模板）
    """
    tpl = template_path or str(_TEMPLATE_PATH)
    if not Path(tpl).exists():
        raise FileNotFoundError(f"MSDS 标准模板不存在: {tpl}")

    doc = Document(tpl)

    # 预拆分组合字段
    section_parts: dict[str, dict[str, str]] = {}
    for section, field_key in _SECTION_FIELD.items():
        aliases = {
            a
            for _t in _SUB_LABEL_SPECS
            if _SUB_LABEL_SPECS[_t][0] == section
            for a in _SUB_LABEL_SPECS[_t][1]
        }
        section_parts[section] = _split_into_parts(entry.get(field_key) or "", list(aliases))

    # 遍历表格单元格，填充占位符
    for table in doc.tables:
        for row in table.rows:
            seen: set[int] = set()
            last_label = ""
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                for para in cell.paragraphs:
                    text = para.text
                    if "【" not in text:
                        norm = _normalize(text)
                        if norm in _DIRECT_LABEL_TO_FIELD or norm in _SUB_LABEL_SPECS or norm in _SECTION_HEADINGS:
                            last_label = norm
                        continue
                    label = _normalize(_label_before_placeholder(text))
                    value = _resolve_value(label, last_label, entry, section_parts)
                    _replace_placeholder(para, value)
                    if label:
                        last_label = label

    doc.save(output_path)
    logger.info("MSDS 标准 docx 已填充模板: %s", output_path)
    return output_path


def standard_docx_filename(entry: dict, collection_id: str, idx: int) -> str:
    """标准 docx 文件名（sanitized）。"""
    name = (entry.get("name") or "unknown").replace("/", "_").replace("\\", "_")
    cas = (entry.get("cas_no") or "no_cas").replace("/", "_").replace("\\", "_")
    return f"{name}_{cas}_{collection_id[:8]}_{idx}.docx"
