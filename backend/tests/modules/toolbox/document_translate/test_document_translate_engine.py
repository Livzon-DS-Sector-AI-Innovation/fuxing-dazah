"""vendor 引擎离线自检（mock 模式，不调模型）：docx / pdf 两条流水线 + 进度回调。

引擎自 ai_translator 项目原样拷贝（仅增加 progress_cb），此处验证其在
backend 环境下可完整跑通并产出翻译文件与报告。
"""

import re
import threading
import time
from pathlib import Path
from types import SimpleNamespace
from typing import Any, cast

import pytest
from docx import Document

from app.modules.toolbox.tools.document_translate._engine.cli import run_pipeline
from app.modules.toolbox.tools.document_translate._engine.config import Settings
from app.modules.toolbox.tools.document_translate._engine.glossary import Glossary
from app.modules.toolbox.tools.document_translate._engine.runner import translate_all
from app.modules.toolbox.tools.document_translate._engine.translate import (
    LLMTranslator,
    make_batches,
)


def make_docx(path: Path) -> Path:
    doc = Document()
    doc.add_heading("Equipment Cleaning Procedure", level=1)
    doc.add_paragraph("The operator shall verify the cleanliness of the machine before starting production.")
    doc.add_paragraph("操作人员应在开始生产前确认设备已清洁。")
    doc.save(str(path))
    return path


def make_pdf(path: Path) -> Path:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 90), "Validation Protocol", fontsize=16)
    page.insert_text(
        (72, 130),
        "This document describes the validation approach for the packaging line.",
        fontsize=11,
    )
    doc.save(str(path))
    doc.close()
    return path


def collect_progress() -> tuple[list[tuple[int, str]], Any]:
    seen: list[tuple[int, str]] = []

    def cb(percent: int, message: str) -> None:
        seen.append((percent, message))

    return seen, cb


def assert_progress_sequence(seen: list[tuple[int, str]]) -> None:
    """进度以低百分比开始、100 结束、单调不减，消息非空。"""
    assert seen, "进度回调未被调用"
    percents = [p for p, _ in seen]
    assert percents[0] <= 10
    assert percents[-1] == 100
    assert percents == sorted(percents)
    assert all(m for _, m in seen)


def test_docx_pipeline_mock(tmp_path: Path) -> None:
    src = make_docx(tmp_path / "protocol.docx")
    seen, cb = collect_progress()
    s = Settings(
        input=str(src), output=str(tmp_path / "out.docx"), report=str(tmp_path / "out.report.md"),
        mode="bilingual", direction="auto", glossary="", mock=True,
    )
    rc = run_pipeline(s, cb)

    assert rc == 0
    out = tmp_path / "out.docx"
    report = tmp_path / "out.report.md"
    assert out.exists() and report.exists()
    # 双语对照：输出段落数多于原文（插入了译文行）
    assert len(Document(str(out)).paragraphs) > len(Document(str(src)).paragraphs)
    assert "翻译校验报告" in report.read_text(encoding="utf-8")
    assert_progress_sequence(seen)


def test_docx_pipeline_mock_replace_mode(tmp_path: Path) -> None:
    src = make_docx(tmp_path / "protocol.docx")
    s = Settings(
        input=str(src), output=str(tmp_path / "out.docx"), report=str(tmp_path / "out.report.md"),
        mode="replace", direction="auto", glossary="", mock=True,
    )
    rc = run_pipeline(s, None)

    assert rc == 0
    # 整段替换：段落数与原文一致
    assert len(Document(str(tmp_path / "out.docx")).paragraphs) == len(Document(str(src)).paragraphs)


def test_docx_pipeline_rejects_unknown_suffix(tmp_path: Path) -> None:
    """引擎对 .docx/.pdf 之外的输入直接拒绝（退出码 2、无产物）。"""
    src = tmp_path / "data.xlsx"
    src.write_bytes(b"fake")
    s = Settings(input=str(src), glossary="", mock=True)
    rc = run_pipeline(s, None)
    assert rc == 2


def test_pdf_pipeline_mock(tmp_path: Path) -> None:
    """经 run_pipeline 入口走 PDF 分支——覆盖率较高的直接调用曾漏掉 progress_cb 未透传的 bug。"""
    pytest.importorskip("fitz")
    src = make_pdf(tmp_path / "validation.pdf")
    seen, cb = collect_progress()
    s = Settings(
        input=str(src), output=str(tmp_path / "out.docx"), report=str(tmp_path / "out.report.md"),
        mode="bilingual", direction="auto", glossary="", mock=True,
        profile_pages=0, pages="",
    )
    rc = run_pipeline(s, cb)

    assert rc == 0
    out = tmp_path / "out.docx"
    report = tmp_path / "out.report.md"
    assert out.exists() and report.exists()
    reopened = Document(str(out))
    assert reopened.paragraphs, "PDF 重排输出不应为空文档"
    assert "翻译校验报告" in report.read_text(encoding="utf-8")
    assert_progress_sequence(seen)


# ── 结构化输出降级：只有服务商真的拒绝 json_schema 才全局关闭 ──

class _FakeCompletions:
    """假 completions：按 response_format 类型决定是否抛 400，并记录调用顺序。"""

    def __init__(self, fail_on: set[str]) -> None:
        self.calls: list[str] = []
        self._fail_on = fail_on

    def create(self, **kwargs: Any) -> Any:
        fmt = kwargs["response_format"]["type"]
        self.calls.append(fmt)
        if fmt in self._fail_on:
            raise _bad_request("this is not a schema error")
        return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content='{"t": []}'))])


def _bad_request(message: str) -> Any:
    import httpx
    from openai import BadRequestError

    return BadRequestError(
        message, response=httpx.Response(400, request=httpx.Request("POST", "http://test")), body=None
    )


def _translator_with(fail_on: set[str]) -> tuple[Any, _FakeCompletions]:
    tr = LLMTranslator("http://test", "sk-test", "test-model")
    completions = _FakeCompletions(fail_on)
    tr.client = cast(Any, SimpleNamespace(chat=SimpleNamespace(completions=completions)))
    return tr, completions


def test_non_schema_400_keeps_structured_output() -> None:
    """与 schema 无关的 400（超长/限流）不得关掉结构化输出。

    降级重试同样失败说明问题不在 schema：此时若照样关闭，本轮剩余批次会静默
    失去 schema 约束（旧实现无条件关闭，正是这个用例覆盖的回归）。
    """
    from openai import BadRequestError

    tr, completions = _translator_with({"json_schema", "json_object"})

    with pytest.raises(BadRequestError):
        tr._request("hi")

    assert completions.calls == ["json_schema", "json_object"]
    assert tr._schema_ok is True


def test_schema_rejected_disables_structured_output_after_fallback() -> None:
    """降级成功才判定服务商不支持 json_schema：本批及后续都走 json_object。"""
    tr, completions = _translator_with({"json_schema"})

    tr._request("hi")
    assert completions.calls == ["json_schema", "json_object"]
    assert tr._schema_ok is False

    tr._request("hi again")
    assert completions.calls == ["json_schema", "json_object", "json_object"]


# ── 占位符重译：并发执行 + 尾部进度可见 ──

class _DroppingTranslator:
    """首轮剥掉全部占位符（触发重译），重译时原样返回掩码文本（还原成功）。

    记录重译请求的在途峰值：串行实现下恒为 1，并发实现下应大于 1。
    """

    def __init__(self, delay: float = 0.05) -> None:
        self._lock = threading.Lock()
        self._inflight = 0
        self.max_inflight = 0
        self._delay = delay

    def translate_batch(
        self,
        items: list[tuple[int, str]],
        target_lang: str,
        doc_type: str,
        glossary_lines: list[str],
        prev_context: str,
    ) -> dict[int, str]:
        if "重试提示" not in (prev_context or ""):
            # 首轮：模拟模型漏发占位符
            return {i: re.sub(r"⟦[^⟧]*⟧", "", text) for i, text in items}
        with self._lock:
            self._inflight += 1
            self.max_inflight = max(self.max_inflight, self._inflight)
        try:
            time.sleep(self._delay)  # 拉开窗口
            return {i: text for i, text in items}
        finally:
            with self._lock:
                self._inflight -= 1


class _Flags:
    def __init__(self) -> None:
        self.flags: list[tuple[str, str]] = []

    def add(self, category: str, message: str) -> None:
        self.flags.append((category, message))


def _make_units(n: int) -> list[SimpleNamespace]:
    return [
        SimpleNamespace(
            idx=i, text=f"第 {i} 批物料 25.3 kg 批号 BN2026001，温度 80±5 ℃。",
            lang="zh", in_table=False, translation="",
        )
        for i in range(n)
    ]


@pytest.mark.parametrize("concurrency", [1, 4])
def test_placeholder_retry_is_concurrent_and_reports_tail_progress(concurrency: int) -> None:
    """占位符重译必须并发，且尾部进度可见。

    数字密集型文档攒下十几个异常单元很常见：串行重译会在并发阶段跑完后
    再挂一条几十分钟的尾巴（且占着引擎锁），进度条也会长时间停在批次阶段的末尾。
    """
    units = _make_units(4)
    s = Settings(input="x.docx", direction="en2zh", concurrency=concurrency, mock=False,
                 date_format="keep")
    translator = _DroppingTranslator()
    ctx = _Flags()
    pos_of = {id(u): i for i, u in enumerate(units)}
    progress: list[tuple[float, str]] = []

    failed = translate_all(
        make_batches(units), translator, s, Glossary([]), units, pos_of, ctx,
        progress_cb=lambda frac, msg: progress.append((frac, msg)),
    )

    assert failed == []
    assert all(u.translation for u in units), "重译后译文应已回写"
    assert any(cat == "占位符重试成功" for cat, _ in ctx.flags)
    assert translator.max_inflight == (1 if concurrency == 1 else 4)

    fracs = [f for f, _ in progress]
    assert fracs == sorted(fracs), "进度必须单调不减"
    assert any(m.startswith("重译占位符异常段落") for _, m in progress)
    assert fracs[-1] > 0.9, "重译阶段必须落在阶段进度的尾段（0.9~1.0）"
