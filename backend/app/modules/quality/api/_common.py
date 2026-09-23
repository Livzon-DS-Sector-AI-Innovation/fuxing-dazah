"""Quality 模块 API 路由。"""

import re
import uuid
from io import BytesIO
from pathlib import Path
from urllib.parse import quote

from docx import Document
from fastapi import (
    APIRouter,
)
from fastapi.responses import StreamingResponse

from app.modules.quality import storage as quality_storage
from app.modules.quality.models import QualityStandardDocument
from app.modules.quality.report_generator import (
    extract_template_placeholders,
    format_value,
    parse_placeholder,
    resolve_placeholders,
)
from app.shared.module_registry import MODULES_BY_CODE

router = APIRouter()
_module = MODULES_BY_CODE["quality"]
REPORT_TEMPLATE_DIR = quality_storage.REPORT_TEMPLATE_DIR
PLACEHOLDER_RE = re.compile(r"\{\{(.+?)\}\}")

# ─── 模板管理 helpers ───
def _scan_templates(root: Path, prefix: str = "") -> list[dict]:
    items = []
    for p in sorted(root.iterdir()):
        if p.is_dir() and p.name == "_generated":
            continue  # 生成输出目录，不参与模板扫描
        rel = f"{prefix}/{p.name}" if prefix else p.name
        if p.is_dir():
            items.append({"type": "folder", "name": p.name, "children": _scan_templates(p, rel)})
        elif p.suffix == ".docx":
            phs = extract_template_placeholders(str(p))
            st = p.stat()
            items.append({
                "type": "template",
                "filename": p.name,
                "path": rel,
                "size_kb": round(st.st_size / 1024, 1),
                "modified": st.st_mtime,
                "placeholder_count": len(phs),
                "placeholders": [
                    {"name": ph.name, "decimals": ph.decimals, "suffix": ph.suffix}
                    for ph in phs
                ],
            })
    return items


# ─── 模板 COA 号 ↔ 标准文档号码匹配 ───

_TEMPLATE_TOKEN_RE = re.compile(r"SOP\.\d{2}\.\d{3,4}(?:\.\d{3})?|EX-[A-Z]+-\d+-\d+", re.IGNORECASE)


def _extract_number_tokens(*texts: str) -> list[str]:
    """提取文本中的号码片段（SOP 号、EX 表号、4-6 位数字）。"""
    tokens: list[str] = []
    for t in texts:
        if not t:
            continue
        for m in _TEMPLATE_TOKEN_RE.findall(t):
            tokens.append(m.upper())
        for m in re.findall(r"\d{4,6}", t):
            tokens.append(m)
    return tokens


def _template_number_tokens(path: Path) -> list[str]:
    """模板号码 token：文件名 + 文档内容（COA 号通常写在表头）。"""
    tokens = _extract_number_tokens(path.stem)
    try:
        doc = Document(str(path))
        texts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    texts.append(cell.text)
        tokens += _extract_number_tokens(*texts)
    except Exception:
        pass
    return tokens


def _score_tokens(a: list[str], b: list[str]) -> int:
    score = 0
    for x in a:
        for y in b:
            if x == y:
                score += 10
            elif len(x) >= 3 and len(y) >= 3 and (x in y or y in x):
                score += 3
    return score


def _match_template_to_doc(
    tp_tokens: list[str],
    docs: list[QualityStandardDocument],
    doc_tokens: dict[uuid.UUID, list[str]],
) -> QualityStandardDocument | None:
    """模板号码与标准文档号码匹配：唯一最高分且 >0 返回文档，平手或零分返回 None。"""
    if not tp_tokens:
        return None
    scored = sorted(
        ((d, _score_tokens(tp_tokens, doc_tokens[d.id])) for d in docs if _score_tokens(tp_tokens, doc_tokens[d.id]) > 0),
        key=lambda x: -x[1],
    )
    if not scored:
        return None
    if len(scored) > 1 and scored[0][1] == scored[1][1]:
        return None  # 平手不自动绑定，留人工判断
    return scored[0][0]


def _doc_match_info(d: QualityStandardDocument | None) -> dict | None:
    if not d:
        return None
    return {"doc_id": str(d.id), "file_no": d.file_no, "product_code": d.product_code or ""}

# ─── 产品代码管理 ───

PRODUCTS_FILE = REPORT_TEMPLATE_DIR.parent / "products.json"

# ─── 产品代码 helpers ───
def _load_products() -> list[dict]:
    import json

    # MinIO 兜底：本地丢失时从对象恢复
    if not PRODUCTS_FILE.exists():
        data = quality_storage.read_products_from_minio()
        if data:
            PRODUCTS_FILE.write_bytes(data)
    if PRODUCTS_FILE.exists():
        try:
            return json.loads(PRODUCTS_FILE.read_text())
        except Exception:
            pass
    return []


def _save_products(data: list[dict]):
    import json

    PRODUCTS_FILE.write_text(json.dumps(data, ensure_ascii=False, indent=2))
    quality_storage.upload_products(PRODUCTS_FILE.read_bytes())

# ─── 报告渲染 helpers ───
def _replace_placeholders(para, data: dict):
    full = para.text
    matches = list(PLACEHOLDER_RE.finditer(full))
    if not matches:
        return
    new_text = full
    for m in reversed(matches):
        raw = m.group(1).strip()
        spec = parse_placeholder(raw)
        val = data.get(spec.name)
        text = format_value(val, spec) if val not in (None, "") else "-"
        new_text = new_text[: m.start()] + text + new_text[m.end() :]
    if para.runs:
        para.runs[0].text = new_text
        for r in para.runs[1:]:
            r.text = ""


def _render_cao_file(
    tp: Path, fill_data: dict, product_name: str, batch_number: str, suffix: str = ""
) -> tuple[Path, str, bytes]:
    """填充模板并落盘到 _generated 目录，返回 (输出路径, 文件名, 文件内容)。

    suffix 用于逐份生成时区分标准文件（文件名追加文件编号）。
    """
    tp = quality_storage.ensure_template_local(tp.relative_to(REPORT_TEMPLATE_DIR).as_posix())
    doc = Document(str(tp))
    # 收集模板全部占位符名，做双向解析（简写/编号变体兜底命中数据键）
    names: set[str] = set()
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for m in PLACEHOLDER_RE.finditer(para.text):
                        names.add(parse_placeholder(m.group(1).strip()).name)
    for para in doc.paragraphs:
        for m in PLACEHOLDER_RE.finditer(para.text):
            names.add(parse_placeholder(m.group(1).strip()).name)
    fill_data = resolve_placeholders(fill_data, list(names))
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    _replace_placeholders(para, fill_data)
    for para in doc.paragraphs:
        _replace_placeholders(para, fill_data)

    output_dir = REPORT_TEMPLATE_DIR / "_generated"
    output_dir.mkdir(parents=True, exist_ok=True)
    safe_batch = batch_number.replace("/", "_").replace("\\", "_")
    safe_suffix = suffix.replace("/", "_").replace("\\", "_")
    # 文件名含流水号：同批号重复出报不再同名覆盖（否则两条记录下载同一文件，静默错单）
    safe_serial = str(fill_data.get("流水号") or "").replace("/", "_").replace("\\", "_")
    output_filename = (
        f"COA-{product_name}-{safe_batch}"
        + (f"-{safe_serial}" if safe_serial else "")
        + (f"-{safe_suffix}" if safe_suffix else "")
        + ".docx"
    )
    output_path = output_dir / output_filename
    doc.save(str(output_path))
    quality_storage.upload_report(output_path)
    return output_path, output_filename, output_path.read_bytes()


def _safe_filename(name: str) -> str:
    """客户端文件名消毒：只取最后一段路径并过滤控制字符，防路径穿越（附件 key 用）。"""
    return re.sub(r"[^\w.\-（）()\[\] ]", "_", Path(name).name) or "file"


def _coa_response(output_filename: str, content: bytes) -> StreamingResponse:
    """中文文件名用 RFC 5987 filename* 编码（latin-1 头编不了中文）。"""
    encoded = quote(output_filename)
    return StreamingResponse(
        BytesIO(content),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": f"attachment; filename=\"coa.docx\"; filename*=UTF-8''{encoded}"
        },
    )
